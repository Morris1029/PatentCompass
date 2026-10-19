"""Exercise the real HTTP API, Linux worker and original patent workbook."""

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import time
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
import httpx
import openpyxl


ROOT = Path(__file__).resolve().parents[1]


def check(condition, message):
    if not condition:
        raise RuntimeError(message)


def reference_statistics(path):
    """Calculate source totals independently of the application's importer."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    years, applicants, ipc, publications = Counter(), Counter(), Counter(), set()
    applications = abstracts = 0
    split = lambda value: list(dict.fromkeys(part.strip() for part in re.split(r"[;；\r\n]+", str(value or "")) if part.strip()))
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows)]
            for values in rows:
                row = dict(zip(headers, values))
                if not row.get("申请号"):
                    continue
                applications += 1
                abstracts += bool(str(row.get("摘要") or "").strip())
                value = row["申请日"]
                years[int(value.year if isinstance(value, datetime) else str(value)[:4])] += 1
                applicants.update(split(row["申请（专利权）人"]))
                ipc.update(split(row["IPC分类号"]))
                publications.update(split(row["公开（公告）号"]))
    finally:
        workbook.close()
    return {"application_count": applications, "publication_count": len(publications), "abstract_count": abstracts,
            "yearly": [{"year": year, "count": count} for year, count in sorted(years.items())],
            "applicants": [{"name": name, "count": count} for name, count in sorted(applicants.items(), key=lambda pair: (-pair[1], pair[0]))],
            "ipc": [{"code": code, "count": count} for code, count in sorted(ipc.items(), key=lambda pair: (-pair[1], pair[0]))]}


def run(args):
    sample = args.sample.resolve()
    source_hash = hashlib.sha256(sample.read_bytes()).hexdigest()
    expected = reference_statistics(sample)
    check(tuple(expected[key] for key in ("application_count", "publication_count", "abstract_count")) == (520, 526, 520),
          "The supplied acceptance sample must contain 520 applications, 526 publication identifiers and 520 abstracts")
    identity = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8]
    output = args.output_root.resolve() / identity
    output.mkdir(parents=True)
    summary = {"acceptance_id": identity, "base_url": args.base_url, "sample_sha256": source_hash,
               "expected": expected, "runs": {}, "checks": [], "status": "running"}

    def note(message):
        print(message, flush=True)
        summary["checks"].append(message)

    with httpx.Client(base_url=args.base_url.rstrip("/"), timeout=30) as client:
        def request(method, path, *, expected_status=200, key=None, **kwargs):
            headers = dict(kwargs.pop("headers", {}))
            if method in {"POST", "PUT", "PATCH"}:
                headers["Idempotency-Key"] = f"{identity}-{key or uuid4().hex}"
            response = client.request(method, path, headers=headers, **kwargs)
            check(response.status_code == expected_status,
                  f"{method} {path}: expected {expected_status}, got {response.status_code}: {response.text[:1500]}")
            return response.json()

        def wait_run(run_id, waiting=False):
            deadline = time.monotonic() + args.timeout
            while time.monotonic() < deadline:
                result = request("GET", f"/api/runs/{run_id}")
                if waiting and result["status"] == "waiting_input":
                    return result
                if result["status"] in {"succeeded", "failed", "cancelled"}:
                    check(result["status"] == "succeeded", f"Run {run_id} failed: {result.get('error')}")
                    return result
                time.sleep(0.5)
            raise RuntimeError(f"Run {run_id} did not finish within {args.timeout} seconds")

        try:
            readiness = request("GET", "/health/ready")
            check(readiness["status"] == "ready", "Dependencies are not ready")
            summary["readiness"] = readiness
            openapi = request("GET", "/openapi.json")
            business_operations = sum(method in {"get", "post", "put", "patch", "delete"}
                                      for path, methods in openapi["paths"].items() if path.startswith("/api/") for method in methods)
            check(business_operations == 29, f"Expected 29 business operations, got {business_operations}")
            note("Readiness and 29 business operations verified")

            project_body = {"name": f"无人机样本接口验收 {identity}",
                            "scope": {"industry": "无人机", "source_region": "中国", "coverage": "部分检索结果"}}
            project = request("POST", "/api/projects", json=project_body, expected_status=201, key="project")
            replay = request("POST", "/api/projects", json=project_body, expected_status=201, key="project")
            check(replay["id"] == project["id"], "Idempotent project creation returned a different ID")
            project_id = summary["project_id"] = project["id"]
            note("Project created; idempotent replay verified")

            with sample.open("rb") as file:
                imported = request("POST", "/api/datasets/imports", expected_status=202,
                    data={"name": "无人机中国地区部分导出样本", "source_metadata": json.dumps(
                        {"query": "无人机", "region": "中国", "partial_results": True}, ensure_ascii=False)},
                    files={"file": (sample.name, file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            summary["dataset"] = imported
            summary["runs"]["import"] = wait_run(imported["run_id"])
            quality = request("GET", f"/api/datasets/{imported['dataset_id']}/quality",
                              params={"dataset_version_id": imported["dataset_version_id"]})
            check(quality["status"] == "ready", "Imported dataset is not ready")
            for key in ("application_count", "publication_count", "abstract_count"):
                check(quality["quality"][key] == expected[key], f"Quality mismatch: {key}")
            summary["quality"] = quality
            reused = request("POST", "/api/datasets/imports", expected_status=202,
                             data={"source_file_id": imported["source_file_id"]})
            check(reused["reused"] and reused["dataset_version_id"] == imported["dataset_version_id"], "Repeated import did not reuse dataset")
            raw = client.get(f"/api/files/{imported['source_file_id']}")
            check(raw.status_code == 200 and hashlib.sha256(raw.content).hexdigest() == source_hash, "Stored source differs from original")
            note("Imported 520 applications / 526 publication IDs / 520 abstracts; duplicate import and raw file hash verified")

            request("PUT", f"/api/projects/{project_id}/datasets", json={"expected_revision": project["revision"],
                    "dataset_version_ids": [imported["dataset_version_id"]]})
            analyzing = request("POST", f"/api/projects/{project_id}/runs", expected_status=202, json={"action": "analyze"})
            pending = wait_run(analyzing["run_id"], waiting=True)
            check(pending["status"] == "waiting_input", "Unconfirmed scope should require input")
            request("POST", f"/api/runs/{analyzing['run_id']}/inputs", expected_status=202,
                    json={"input_request_id": pending["input_request"]["id"],
                          "expected_scope_version_id": pending["scope_version_id"], "values": {"confirmed": True}})
            analyzed = summary["runs"]["analyze"] = wait_run(analyzing["run_id"])
            analysis = request("GET", f"/api/projects/{project_id}/analysis", params={"run_id": analyzed["id"]})
            for key, value in expected.items():
                check(analysis["summary"][key] == value, f"Independent source statistics differ: {key}")
            summary["analysis"] = analysis
            note("Scope input resumed the same run; all statistics match independent workbook calculations")

            patents = []
            cursor = None
            while True:
                page = request("GET", f"/api/projects/{project_id}/patents", params={"limit": 200, **({"cursor": cursor} if cursor else {})})
                patents.extend(page["items"])
                cursor = page["next_cursor"]
                if not cursor:
                    break
            check(len(patents) == len({item["id"] for item in patents}) == 520, "Patent pagination omitted or repeated records")
            chart_records = []
            for chart_id in analyzed["result"]["chart_ids"]:
                chart = request("GET", f"/api/charts/{chart_id}")
                check(chart["provenance"]["dataset_version_ids"] == [imported["dataset_version_id"]], "Chart dataset snapshot differs")
                check(len(chart["provenance"]["application_ids"]) == 520, "Chart omitted source applications")
                dimension = chart["data"]["rows"][0]["value"]
                down = request("GET", f"/api/projects/{project_id}/patents", params={"chart_id": chart_id,
                    "chart_version_id": chart["version_id"], "dimension_value": dimension, "limit": 200})
                check(down["total"] == chart["data"]["rows"][0]["count"], "Chart drilldown differs from saved count")
                chart_records.append(chart)
            check(len(chart_records) == 3, "Expected annual, applicant and IPC charts")
            evidence = [request("GET", f"/api/evidence/{evidence_id}") for evidence_id in analyzed["result"]["evidence_ids"]]
            check(any(item["kind"] == "patent" and item["content"]["abstract"] for item in evidence), "Original abstract evidence missing")
            check(sum(item["kind"] == "statistic" for item in evidence) == 3, "Statistical evidence missing")
            summary["charts"], summary["evidence"] = chart_records, evidence
            events = client.get(f"/api/runs/{analyzed['id']}/events")
            check(events.status_code == 200 and "event: run.completed" in events.text, "Persisted SSE replay is incomplete")
            note("Patent pagination, three charts, snapshot drilldown, original evidence and SSE replay verified")

            generated = request("POST", f"/api/projects/{project_id}/runs", expected_status=202,
                                json={"action": "generate_report", "analysis_id": analyzed["result"]["analysis_id"], "report_mode": "statistical"})
            report_run = summary["runs"]["report"] = wait_run(generated["run_id"])
            report_id = summary["report_id"] = report_run["result"]["report_id"]
            v1 = request("GET", f"/api/reports/{report_id}/versions/{report_run['result']['report_version_id']}")
            check(v1["number"] == 1, "New report did not start at version 1")
            block = next(item for item in v1["content"]["blocks"] if item["type"] == "paragraph")

            def revise(base, suffix, label):
                task = request("POST", f"/api/reports/{report_id}/revisions", expected_status=202,
                    json={"base_version_id": base["id"], "expected_current_version_id": base["id"],
                          "block_id": block["id"], "replacement_text": block["text"] + suffix})
                completed = summary["runs"][label] = wait_run(task["run_id"])
                return request("GET", f"/api/reports/{report_id}/versions/{completed['result']['report_version_id']}")

            marker_v2, marker_v3 = " HTTP_ACCEPTANCE_V2", " HTTP_ACCEPTANCE_V3"
            v2 = revise(v1, marker_v2, "revision_v2")
            check(v2["number"] == 2 and v2["parent_id"] == v1["id"], "Version 2 lineage is wrong")
            request("PUT", f"/api/reports/{report_id}/current-version", json={"version_id": v1["id"], "expected_current_version_id": v2["id"]})
            v3 = revise(v1, marker_v3, "revision_v3")
            check(v3["number"] == 3 and v3["parent_id"] == v1["id"], "Branch revision did not preserve parent version 1")
            check(request("GET", f"/api/reports/{report_id}/versions/{v1['id']}") == v1, "Original version 1 was overwritten")
            check(request("GET", f"/api/reports/{report_id}/versions/{v2['id']}") == v2, "Version 2 branch was overwritten")
            versions = request("GET", f"/api/reports/{report_id}/versions")
            check({item["number"] for item in versions["items"]} == {1, 2, 3}, "Report version history is incomplete")
            project = request("GET", f"/api/projects/{project_id}")
            check(project["current_report_version_id"] == v3["id"], "Current adopted version is not version 3")
            summary["versions"] = {"v1": v1, "v2": v2, "v3": v3}
            note("Report V1 -> V2 -> adopt V1 -> V3 verified; original and sibling branch remain immutable")

            task = request("POST", f"/api/reports/{report_id}/exports", expected_status=202,
                           json={"report_version_id": v1["id"], "format": "docx"})
            exported = summary["runs"]["export_v1"] = wait_run(task["run_id"])
            check(exported["result"]["report_version_id"] == v1["id"], "Export used current version instead of requested version 1")
            document = client.get(f"/api/files/{exported['result']['file_id']}")
            check(document.status_code == 200, "Word download failed")
            docx_path = output / "patentcompass-report-v1.docx"
            docx_path.write_bytes(document.content)
            word = Document(BytesIO(document.content))
            text = "\n".join(paragraph.text for paragraph in word.paragraphs)
            check(block["text"] in text and marker_v2 not in text and marker_v3 not in text, "Word content does not match frozen version 1")
            check("520" in text and "526" in text, "Word source totals differ")
            with ZipFile(BytesIO(document.content)) as archive:
                check(len([name for name in archive.namelist() if name.startswith("word/media/")]) == 3, "Word should embed all three rendered charts")
            check(hashlib.sha256(sample.read_bytes()).hexdigest() == source_hash, "Original workbook was changed")
            summary["export"] = {"path": str(docx_path), "bytes": len(document.content), "sha256": hashlib.sha256(document.content).hexdigest(),
                                 "file_id": exported["result"]["file_id"], "report_version_id": v1["id"]}
            summary["status"] = "passed"
            note("Requested V1 Word downloaded with three images and matching frozen text; source workbook preserved")
        except Exception as exc:
            summary["status"], summary["error"] = "failed", str(exc)
            raise
        finally:
            (output / "result.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"Acceptance result: {output / 'result.json'}", flush=True)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--sample", type=Path, default=ROOT / "database" / "patent-demo.xlsx")
    parser.add_argument("--output-root", type=Path, default=ROOT / "storage" / "acceptance")
    parser.add_argument("--timeout", type=float, default=240, help="Maximum seconds per asynchronous task")
    run(parser.parse_args())
