"""Real HTTP acceptance for materials, merge snapshots and report operations.

The optional oneoff executor loads current mounted source without interrupting a
long-running worker. Every business operation is still created through HTTP and
persists in the same database; normal Celery acceptance is a separate check.
"""

import argparse
from datetime import datetime, timedelta, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import time
from uuid import UUID, uuid4
from zipfile import ZipFile

from docx import Document
import httpx
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def check(value, message):
    if not value:
        raise RuntimeError(message)


def run(args):
    identity = "frontend-data-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8]
    output = ROOT / "storage" / "acceptance" / identity
    output.mkdir(parents=True)
    result = {"identity": identity, "started_at": datetime.now(timezone.utc).isoformat(), "status": "running",
              "executor": args.executor, "stages": {}, "runs": {}}
    if args.resume:
        previous = json.loads(args.resume.read_text(encoding="utf-8"))
        result = {**previous, "identity": identity, "status": "running", "executor": args.executor,
                  "resumed_from": str(args.resume.resolve()), "started_at": datetime.now(timezone.utc).isoformat()}
        result.pop("error", None)
        if args.refresh_export:
            result["stages"].pop("frozen_word_download", None)

    def save():
        (output / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    with httpx.Client(base_url=args.base_url, timeout=45, trust_env=False) as client:
        def request(method, path, expected=200, **kwargs):
            headers = kwargs.pop("headers", {})
            if method in {"POST", "PUT", "PATCH"}:
                headers.setdefault("Idempotency-Key", identity + "-" + uuid4().hex)
            response = client.request(method, path, headers=headers, **kwargs)
            check(response.status_code == expected, f"{method} {path} returned {response.status_code}: {response.text[:1400]}")
            return response.json()

        def execute_fresh(run_id):
            if args.executor == "celery":
                return
            # Reserve only this HTTP-created acceptance run for a fresh process;
            # duplicate deliveries remain protected by the normal lease claim.
            sys.path.insert(0, str(ROOT / "api"))
            from app.db import session_scope
            from app.models import ResearchRun
            with session_scope() as session:
                row = session.get(ResearchRun, str(UUID(run_id)), with_for_update=True)
                if row.status != "queued":
                    return
                row.dispatch_pending = False
                row.dispatched_at = datetime.now(timezone.utc) + timedelta(minutes=15)
            code = f"from app.jobs import execute_run;execute_run('{UUID(run_id)}')"
            command = ["docker", "compose", "-f", "deploy/compose.yml"]
            command += (["run", "--rm", "--no-deps", "worker"] if args.executor == "oneoff" else ["exec", "-T", "worker"])
            process = subprocess.run([*command, "/workspace/api/.venv/bin/python", "-c", code], cwd=ROOT,
                                     capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=args.timeout)
            check(process.returncode == 0, f"Fresh task executor failed for {run_id}: {process.stderr[-1200:]}")

        def wait(task, name, waiting=False):
            run_id = task["run_id"]
            execute_fresh(run_id)
            deadline = time.monotonic() + args.timeout
            while time.monotonic() < deadline:
                value = request("GET", f"/api/runs/{run_id}")
                result["runs"][name] = value
                save()
                if value["status"] in {"succeeded", "failed", "cancelled", "waiting_input"}:
                    check(value["status"] == ("waiting_input" if waiting else "succeeded"),
                          f"{name} {value['status']}: {value.get('error') or value.get('input_request')}")
                    return value
                time.sleep(0.5)
            raise RuntimeError(f"Timeout waiting for {name}: {run_id}")

        def stage(name, callback):
            if args.resume and result["stages"].get(name, {}).get("status") == "passed":
                print("Reusing " + name, flush=True)
                return result["stages"][name]["result"]
            print("Starting " + name, flush=True)
            started = time.monotonic()
            try:
                value = callback()
                result["stages"][name] = {"status": "passed", "seconds": round(time.monotonic() - started, 2), "result": value}
                print("Passed " + name, flush=True)
                save()
                return value
            except Exception as exc:
                result["stages"][name] = {"status": "failed", "error": str(exc)}
                raise

        try:
            def prepare():
                datasets = request("GET", "/api/datasets", params={"limit": 200})["items"]
                selected = next((version for dataset in datasets for version in dataset["versions"]
                                 if version["status"] == "ready" and not version.get("base_version_id") and
                                 tuple(version["quality"].get(key) for key in ("application_count", "publication_count", "abstract_count")) == (520, 526, 520)), None)
                check(selected, "Original ready 520/526/520 sample not found")
                result["base_version"] = selected
                project = request("POST", "/api/projects", expected=201, json={"name": identity,
                    "scope": {"industry": "无人机", "coverage": "部分检索结果", "source_region": "中国"}})
                result["project_id"] = project["id"]
                request("PUT", f"/api/projects/{project['id']}/datasets", json={"expected_revision": project["revision"],
                    "dataset_version_ids": [selected["id"]]})
                request("PATCH", f"/api/projects/{project['id']}/scope", json={"expected_scope_version_id": project["scope_version_id"],
                    "changes": {}, "confirmed": True})
                return {"project_id": project["id"], "dataset_version_id": selected["id"]}
            stage("prepare", prepare)
            project_id = result["project_id"]

            def materials():
                document = Document()
                document.add_paragraph("本次研究需求：核查无人机视觉降落与障碍感知的样本证据。")
                table = document.add_table(rows=1, cols=2)
                table.cell(0, 0).text, table.cell(0, 1).text = "输出目标", "统计与路线图；不证明技术空白"
                document.add_paragraph("请用一句中文概括此资料的研究需求。")
                target = BytesIO()
                document.save(target)
                content = target.getvalue()
                (output / "research-needs.docx").write_bytes(content)
                task = request("POST", f"/api/projects/{project_id}/materials", expected=202,
                    files={"file": ("research-needs.docx", content)})
                wait(task, "parse_material")
                material = request("GET", f"/api/materials/{task['material_id']}")
                check(material["status"] == "ready" and len(material["blocks"]) == 4, "Material blocks were not persisted")
                check(material["blocks"][1]["location"] == {"table": 1, "row": 1, "cell": 1}, "Cell location mismatch")
                downloaded = client.get(material["download_url"])
                check(hashlib.sha256(downloaded.content).hexdigest() == hashlib.sha256(content).hexdigest(), "Original Word file changed")
                task2 = request("POST", f"/api/projects/{project_id}/materials", expected=202,
                    files={"file": ("research-needs.docx", content)})
                check(task2["material_id"] == material["id"] and task2["reused"], "Duplicate material created a new identity")
                old_chat = result["runs"].get("attachment_only_chat") if args.resume else None
                task = (request("POST", f"/api/conversations/{result['conversation_id']}/messages", expected=202,
                            json={"content": "", "attachments": [material["id"]], "action": "chat"}) if old_chat and old_chat["status"] == "failed" else
                        request("POST", f"/api/projects/{project_id}/conversations", expected=202,
                            json={"initial_message": {"content": "", "attachments": [material["id"]], "action": "chat"}}))
                result["conversation_id"] = task["conversation_id"]
                chat = wait(task, "attachment_only_chat")
                messages = request("GET", f"/api/conversations/{task['conversation_id']}/messages")["items"]
                check(any(item["role"] == "user" and item["content"] == "" and item["attachments"] == [material["id"]] for item in messages), "Attachment-only message lost")
                check(any(item["id"] == chat["result"]["message_id"] for item in messages), "Attachment answer not persisted")
                return {"material": material, "messages": messages}
            stage("materials_and_attachment_only", materials)

            if args.include_doc:
                def legacy_doc():
                    material = result["stages"]["materials_and_attachment_only"]["result"]["material"]
                    source = client.get(material["download_url"])
                    check(source.status_code == 200, "Original DOCX unavailable")
                    (output / "legacy-needs.docx").write_bytes(source.content)
                    process = subprocess.run(["docker", "run", "--rm", "--network", "none", "--user", "10001:10001",
                        "-v", str(output) + ":/data", "patentcompass-worker:dev", "libreoffice",
                        "-env:UserInstallation=file:///tmp/acceptance-profile", "--headless", "--convert-to", "doc",
                        "--outdir", "/data", "/data/legacy-needs.docx"], cwd=ROOT, capture_output=True,
                        text=True, encoding="utf-8", errors="replace", timeout=120)
                    path = output / "legacy-needs.doc"
                    check(process.returncode == 0 and path.is_file(), "Real legacy Word fixture conversion failed")
                    content = path.read_bytes()
                    task = request("POST", f"/api/projects/{project_id}/materials", expected=202,
                        files={"file": (path.name, content)})
                    if task["status"] == "failed":
                        task = request("POST", f"/api/runs/{task['run_id']}/retry", expected=202)
                    wait(task, "legacy_doc_parse")
                    material_id = task.get("material_id") or result["runs"]["legacy_doc_parse"].get("result", {}).get("material_id")
                    converted = request("GET", f"/api/materials/{material_id}")
                    check([block["text"] for block in converted["blocks"]] == [block["text"] for block in material["blocks"]],
                          "Legacy conversion changed extracted body/table text")
                    check(all("page" not in block["location"] for block in converted["blocks"]), "Legacy parser fabricated page numbers")
                    original = client.get(converted["download_url"]).content
                    check(hashlib.sha256(original).hexdigest() == hashlib.sha256(content).hexdigest(), "Legacy original was modified")
                    return {"material_id": converted["id"], "block_count": len(converted["blocks"]), "sha256": hashlib.sha256(original).hexdigest()}
                stage("real_legacy_doc_conversion", legacy_doc)

            def initial_report():
                analysis = wait(request("POST", f"/api/projects/{project_id}/runs", expected=202,
                    json={"action": "analyze", "semantic_mode": "off"}), "sample_statistics")
                report = wait(request("POST", f"/api/projects/{project_id}/runs", expected=202,
                    json={"action": "generate_report", "analysis_id": analysis["result"]["analysis_id"]}), "initial_report")
                result["report_id"] = report["result"]["report_id"]
                version = request("GET", f"/api/reports/{result['report_id']}/versions/{report['result']['report_version_id']}")
                check(version["content"]["source"]["summary"]["application_count"] == 520, "Initial report lost original sample count")
                result["v1"] = version
                return {"report_id": result["report_id"], "version_id": version["id"]}
            stage("original_sample_report", initial_report)

            def report_operations():
                base = result["v1"]
                texts = [block for block in base["content"]["blocks"] if block["type"] == "paragraph"][:2]
                chart = next(block for block in base["content"]["blocks"] if block["type"] == "chart")
                operations = [{"op": "set_title", "text": "无人机接口验收报告"},
                    {"op": "replace_text", "block_id": texts[0]["id"], "text": "本报告固定原始样本：520 条申请、526 个公开标识、520 份摘要。"},
                    {"op": "replace_text", "block_id": texts[1]["id"], "text": "本次仅覆盖已导入部分样本，不能代表全部检索结果。"},
                    {"op": "insert_block", "after_id": texts[1]["id"], "block": {"id": "acceptance-chapter", "type": "heading", "text": "验收说明", "level": 1}},
                    {"op": "insert_block", "after_id": "acceptance-chapter", "block": {"id": "acceptance-paragraph", "type": "paragraph", "text": "补充数据和新范围不会覆盖本历史版本。"}},
                    {"op": "move_block", "block_id": "acceptance-paragraph", "after_id": "acceptance-chapter"},
                    {"op": "set_chart_layout", "block_id": chart["id"], "layout": {"width_percent": 50, "alignment": "right"}}]
                task = request("POST", f"/api/reports/{result['report_id']}/revisions", expected=202,
                    json={"base_version_id": base["id"], "expected_current_version_id": base["id"], "operations": operations,
                          "conversation_id": result["conversation_id"]})
                run = wait(task, "report_operations")
                version = request("GET", f"/api/reports/{result['report_id']}/versions/{run['result']['report_version_id']}")
                check(version["number"] == 2 and version["content"]["title"] == "无人机接口验收报告", "Report operations did not create version 2")
                check(request("GET", f"/api/reports/{result['report_id']}/versions/{base['id']}") == base, "Old report changed")
                check(version["conversation_id"] == result["conversation_id"], "Source conversation was lost")
                result["v2"] = version
                return {"version_id": version["id"], "operations": operations}
            stage("atomic_report_operations", report_operations)

            def merge():
                existing = request("GET", f"/api/projects/{project_id}/patents", params={"limit": 1})["items"][0]
                marker = "validation" + uuid4().hex[:12]
                result["semantic_marker"] = marker
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.append(["申请号", "发明名称", "摘要", "申请日", "公开（公告）号", "公开（公告）日", "申请（专利权）人", "IPC分类号"])
                sheet.append([existing["application_number"], "Existing UAV evidence " + marker, None, None, None, None, "Acceptance supplemental applicant", None])
                fixture_number = "WO2099CN" + str(int(uuid4().hex[:10], 16))
                fixture_publication = "WO2099" + str(int(uuid4().hex[:10], 16)) + "A1"
                sheet.append([fixture_number, "UAV visual landing research " + marker,
                    "A drone uses a visual sensor to identify a landing surface and avoid obstacles during autonomous landing.",
                    "2026.10.19", fixture_publication, "2026.10.19", "Acceptance fixture applicant", "B64U10/13"])
                target = BytesIO()
                workbook.save(target)
                content = target.getvalue()
                (output / "supplement.xlsx").write_bytes(content)
                upload = {"files": {"file": ("supplement.xlsx", content)}, "data": {"base_dataset_version_id": result["base_version"]["id"]}}
                task = request("POST", "/api/datasets/imports", expected=202, **upload)
                waiting = wait(task, "merge_wait", waiting=True)
                quality = request("GET", f"/api/datasets/{task['dataset_id']}/quality", params={"dataset_version_id": task["dataset_version_id"]})
                check(quality["status"] == "waiting_input" and len(quality["conflicts"]) == 2, "Expected title/applicant conflicts were not exposed")
                project = request("GET", f"/api/projects/{project_id}")
                rejected = client.put(f"/api/projects/{project_id}/datasets", json={"expected_revision": project["revision"], "dataset_version_ids": [task["dataset_version_id"]]})
                check(rejected.status_code == 409, "Unresolved merge was bindable")
                resumed = request("POST", f"/api/runs/{task['run_id']}/inputs", expected=202, json={
                    "input_request_id": waiting["input_request"]["id"], "values": {"choices": [
                        {"conflict_id": item["id"], "choice": "merge" if item["field"] == "applicants" else "new"} for item in quality["conflicts"]]}})
                check(resumed["run_id"] == task["run_id"], "Conflict confirmation replaced original run")
                completed = wait(resumed, "merge_ready")
                check(tuple(completed["result"][key] for key in ("application_count", "publication_count", "abstract_count")) == (521, 527, 521), "Merged totals mismatch")
                repeated = request("POST", "/api/datasets/imports", expected=202, **upload)
                check(repeated["dataset_version_id"] == task["dataset_version_id"] and repeated["run_id"] == task["run_id"] and repeated["reused"], "Repeated merge created new data")
                before = request("GET", f"/api/projects/{project_id}/patents", params={"limit": 1})
                check(before["total"] == 520, "Import implicitly adopted new dataset")
                request("PUT", f"/api/projects/{project_id}/datasets", json={"expected_revision": project["revision"], "dataset_version_ids": [task["dataset_version_id"]]})
                check(request("GET", f"/api/projects/{project_id}/patents", params={"limit": 1})["total"] == 521, "Explicit dataset adoption failed")
                old_quality = request("GET", f"/api/datasets/{task['dataset_id']}/quality", params={"dataset_version_id": result["base_version"]["id"]})
                check(tuple(old_quality["quality"][key] for key in ("application_count", "publication_count", "abstract_count")) == (520, 526, 520), "Original dataset totals changed")
                result["merge_version_id"] = task["dataset_version_id"]
                return {"dataset_version_id": task["dataset_version_id"], "conflicts": quality["conflicts"], "result": completed["result"]}
            stage("incremental_merge_conflicts_idempotency", merge)

            if not args.skip_semantic:
                def semantic_and_roadmap():
                    project = request("GET", f"/api/projects/{project_id}")
                    request("PATCH", f"/api/projects/{project_id}/scope", json={"expected_scope_version_id": project["scope_version_id"],
                        "changes": {"q": result["semantic_marker"]}, "confirmed": True})
                    check(request("GET", f"/api/projects/{project_id}/patents", params={"limit": 5})["total"] == 2, "Acceptance semantic scope should contain two records")
                    full = wait(request("POST", f"/api/projects/{project_id}/runs", expected=202,
                        json={"action": "analyze", "semantic_mode": "full"}), "two_record_semantic")
                    report = wait(request("POST", f"/api/projects/{project_id}/runs", expected=202,
                        json={"action": "generate_report", "analysis_id": full["result"]["analysis_id"]}), "semantic_report")
                    version = request("GET", f"/api/reports/{result['report_id']}/versions/{report['result']['report_version_id']}")
                    check(version["content"]["source"]["summary"]["application_count"] == 2, "Semantic report scope mismatch")
                    blocks = version["content"]["blocks"]
                    check(any(block.get("section") == "roadmap" for block in blocks), "Real roadmap absent from semantic report")
                    start = next(index for index, block in enumerate(blocks) if block.get("section") == "roadmap")
                    end = next((index for index in range(start + 1, len(blocks)) if blocks[index]["type"] == "heading" and blocks[index].get("level") == 1), len(blocks))
                    operations = [{"op": "remove_block", "block_id": block["id"]} for block in blocks[start:end]] + [{"op": "append_roadmap"}]
                    changed = wait(request("POST", f"/api/reports/{result['report_id']}/revisions", expected=202, json={
                        "base_version_id": version["id"], "expected_current_version_id": version["id"], "operations": operations}), "append_real_roadmap")
                    result["semantic_report_version_id"] = changed["result"]["report_version_id"]
                    check(request("GET", f"/api/reports/{result['report_id']}/versions/{result['v1']['id']}") == result["v1"], "New semantic scope altered historical report")
                    return {"analysis": full["result"], "version_id": result["semantic_report_version_id"]}
                stage("real_semantic_roadmap_frozen_scope", semantic_and_roadmap)

            def export():
                version = result["v2"]
                exported = wait(request("POST", f"/api/reports/{result['report_id']}/exports", expected=202,
                    json={"report_version_id": version["id"]}), "frozen_v2_export")
                response = client.get(exported["result"]["download_url"])
                check(response.status_code == 200 and response.content.startswith(b"PK"), "Word file not downloadable")
                path = output / "report-v2.docx"
                path.write_bytes(response.content)
                document = Document(BytesIO(response.content))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                check("无人机接口验收报告" in text and "520 条申请、526 个公开标识、520 份摘要" in text, "Frozen Word text/version mismatch")
                check("补充数据和新范围不会覆盖本历史版本" in text, "Inserted chapter text missing")
                check(round(document.inline_shapes[0].width.inches, 2) == 3.25, "Word ignored chart 50% width")
                with ZipFile(BytesIO(response.content)) as archive:
                    images = [name for name in archive.namelist() if name.startswith("word/media/")]
                    check(len(images) == 3, "Expected three chart PNGs")
                result["word_path"] = str(path)
                return {"path": str(path), "sha256": hashlib.sha256(response.content).hexdigest(), "chart_count": len(images), "specified_version_id": version["id"]}
            stage("frozen_word_download", export)
            result["status"] = "passed"
        except Exception as exc:
            result["status"], result["error"] = "failed", str(exc)
            print("Failed: " + str(exc), flush=True)
        finally:
            result["finished_at"] = datetime.now(timezone.utc).isoformat()
            save()
            print(str(output / "result.json"), flush=True)
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--timeout", type=float, default=600)
    parser.add_argument("--executor", choices=["celery", "exec", "oneoff"], default="celery")
    parser.add_argument("--skip-semantic", action="store_true")
    parser.add_argument("--resume", type=Path)
    parser.add_argument("--include-doc", action="store_true")
    parser.add_argument("--refresh-export", action="store_true")
    raise SystemExit(run(parser.parse_args()))
