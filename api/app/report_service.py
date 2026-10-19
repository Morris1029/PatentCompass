"""Immutable, evidence-backed statistical reports and Word exports."""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
from tempfile import TemporaryDirectory
from uuid import NAMESPACE_URL, uuid5

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sqlalchemy import select

from app.db import session_scope
from app.errors import AppError
from app.files import store_bytes
from app.models import (
    AnalysisResult, Chart, DatasetVersion, Evidence, PatentApplication,
    PatentPublication, Project, Report, ReportVersion, ScopeVersion, StoredFile,
)


DEFAULT_REPORT_LAYOUT = {
    "template_version": "statistical-brief-v1", "paper_size": "A4",
    "page_width_inches": 8.27, "page_height_inches": 11.69,
    "margins_inches": {"top": 0.75, "bottom": 0.75, "left": 0.8, "right": 0.8},
    "font_family": "Noto Sans CJK SC", "body_font_pt": 10.5, "chart_width_inches": 6.5,
}


def _version_result(report, version):
    return {"report_id": report.id, "report_version_id": version.id,
            "number": version.number, "current_version_id": report.current_version_id}


def _project(session, project_id):
    project = session.scalar(select(Project).where(Project.id == project_id).with_for_update())
    if not project:
        raise AppError(404, "PROJECT_NOT_FOUND", "项目不存在")
    return project


def _load_snapshot(ctx):
    with session_scope() as session:
        project = session.get(Project, ctx.project_id)
        if not project:
            raise AppError(404, "PROJECT_NOT_FOUND", "项目不存在")
        analysis_id = ctx.payload.get("analysis_id") or project.current_analysis_id
        analysis = session.get(AnalysisResult, analysis_id) if analysis_id else None
        if not analysis or analysis.project_id != ctx.project_id:
            raise AppError(409, "ANALYSIS_REQUIRED", "请先完成本项目的数据分析")
        if (analysis.scope_version_id != ctx.scope_version_id
                or sorted(analysis.dataset_version_ids) != sorted(ctx.dataset_version_ids)):
            raise AppError(409, "ANALYSIS_SNAPSHOT_MISMATCH", "分析结果与任务的数据或范围版本不一致")
        scope = session.get(ScopeVersion, analysis.scope_version_id)
        sources = []
        for version_id in analysis.dataset_version_ids:
            dataset = session.get(DatasetVersion, version_id)
            source_file = session.get(StoredFile, dataset.source_file_id)
            sources.append({"dataset_version_id": version_id,
                            "source_file_id": dataset.source_file_id,
                            "file_name": source_file.original_name, "sha256": source_file.sha256,
                            "rule_version": dataset.rule_version, "quality": dataset.quality,
                            "source_metadata": dataset.source_metadata})
        charts = [{"id": chart.id, "version_id": chart.version_id, "kind": chart.kind,
                   "option": chart.option, "data": chart.data, "provenance": chart.provenance,
                   "selections": chart.selections}
                  for chart in session.scalars(select(Chart).where(
                      Chart.analysis_id == analysis.id).order_by(Chart.kind, Chart.id))]
        evidence = [{"id": row.id, "kind": row.kind, "application_id": row.application_id,
                     "content": row.content}
                    for row in session.scalars(select(Evidence).where(
                        Evidence.analysis_id == analysis.id).order_by(Evidence.id))]
        records = list(session.scalars(select(PatentApplication).where(
            PatentApplication.id.in_(analysis.patent_ids)).order_by(
                PatentApplication.application_date.desc(), PatentApplication.application_number).limit(5)))
        representatives = []
        for row in records:
            publications = list(session.scalars(select(PatentPublication.publication_number).where(
                PatentPublication.application_id == row.id).order_by(PatentPublication.publication_number)))
            representatives.append({"application_id": row.id, "application_number": row.application_number,
                                    "publication_numbers": publications, "title": row.title,
                                    "abstract": row.abstract, "applicants": row.applicants,
                                    "source_sheet": row.source_sheet, "source_row": row.source_row,
                                    "dataset_version_id": row.dataset_version_id})
        return deepcopy({"project_name": project.name, "analysis_id": analysis.id,
                         "analysis_run_id": analysis.run_id, "scope_version_id": analysis.scope_version_id,
                         "scope": scope.content, "dataset_version_ids": analysis.dataset_version_ids,
                         "summary": analysis.summary, "views": analysis.views, "sources": sources,
                         "charts": charts, "evidence": evidence, "representatives": representatives})


def build_report_content(snapshot):
    """Build only facts from the frozen analysis; block identities survive revisions."""
    summary = snapshot["summary"]
    blocks = []

    def add(key, kind, **values):
        block_id = str(uuid5(NAMESPACE_URL, f"patentcompass:{snapshot['analysis_id']}:{key}"))
        blocks.append({"id": block_id, "type": kind, **values})

    add("scope-heading", "heading", text="研究范围与数据说明", level=1)
    add("overview", "paragraph", text=(
        f"本报告统计当前确认范围内的 {summary.get('application_count', 0)} 条申请记录，"
        f"涉及 {summary.get('publication_count', 0)} 个公开标识及 "
        f"{summary.get('abstract_count', 0)} 份行级摘要。结论仅适用于已导入样本，不能代表全部检索结果。"))
    scope_labels = {"query": "检索词", "keyword": "检索词", "region": "地区", "applicant": "申请人",
                    "ipc": "IPC", "application_start": "申请起始日期", "application_end": "申请结束日期"}
    scope_values = {**snapshot["scope"], **(snapshot["scope"].get("filters") or {})}
    scope_description = "；".join(f"{label}：{scope_values[key]}" for key, label in scope_labels.items()
                                  if scope_values.get(key))
    add("scope", "paragraph", text="确认范围：" + (scope_description or "采用本次已确认范围内的全部已导入记录"))
    for index, source in enumerate(snapshot["sources"]):
        metadata, quality = source["source_metadata"], source["quality"]
        conditions = "；".join(f"{label}：{metadata[key]}" for key, label in (
            ("query", "原始检索式"), ("search_query", "原始检索式"), ("region", "原始地区筛选"))
                               if metadata.get(key))
        add(f"source-{index}", "paragraph", text=(
            f"原始文件：{source['file_name']}。{conditions + '。' if conditions else ''}"
            f"导入校验得到 {quality.get('application_count', 0)} 条有效申请、"
            f"{quality.get('publication_count', 0)} 个公开标识，"
            f"记录异常 {len(quality.get('errors', []))} 项。"))
        for warning_index, warning in enumerate(quality.get("warnings", [])):
            add(f"source-{index}-warning-{warning_index}", "paragraph", text=str(warning))
    add("counting", "paragraph", text=(
        "统计以申请记录为单位，共同申请人分别全计数，多 IPC 标签分别计数。"
        "因此申请人或 IPC 合计可超过申请总量。公开标识与申请记录分开统计。"))
    for key, heading, field, columns in (
        ("yearly", "年度申请分布", "yearly", ["年度", "申请数量"]),
        ("applicants", "主要申请人", "applicants", ["申请人", "申请数量"]),
        ("ipc", "IPC 技术分类", "ipc", ["IPC 分类", "申请数量"]),
    ):
        add(f"{key}-heading", "heading", text=heading, level=1)
        rows = summary.get(field, [])
        label = {"yearly": "year", "applicants": "name", "ipc": "code"}[key]
        limit = len(rows) if key == "yearly" else 15
        add(f"{key}-table", "table", columns=columns,
            rows=[[str(row.get(label, "")), str(row["count"])] for row in rows[:limit]],
            caption="全部入选记录统计" if limit == len(rows) else f"按数量展示前 {min(limit, len(rows))} 项")
        for chart in snapshot["charts"]:
            if chart["kind"] == key:
                add(f"{key}-chart", "chart", chart=deepcopy(chart),
                    chart_id=chart["id"], chart_version_id=chart["version_id"],
                    caption=chart["data"].get("title", heading),
                    layout={"width_percent": 100, "alignment": "center"})
    add("limits-heading", "heading", text="可用结论与待补充资料", level=1)
    add("limits", "paragraph", text=(
        "本次提供年度、申请人和 IPC 描述性统计，以及原文摘要定位。"
        "IPC 分类不能直接替代语义技术路线。语义技术路线、机会判断和地域保护布局尚未形成结论；"
        "需要相应模型能力及更完整的原始字段、权利要求和外部补检证据。"))
    add("representatives-heading", "heading", text="样本记录与原文证据附录", level=1)
    add("representatives-method", "paragraph", text=(
        "以下按申请日期从新到旧列出最多五条样本记录，仅用于追溯原文，不代表技术重要性排名。"
        "保留英文摘要原文，不据此推断授权或法律状态。"))
    for record in snapshot["representatives"]:
        key = record["application_id"]
        add(f"patent-{key}", "heading", level=2,
            text=f"{record['application_number']} {record['title']}")
        add(f"patent-source-{key}", "paragraph", text=(
            f"公开编号：{', '.join(record['publication_numbers'])}；申请人：{'; '.join(record['applicants'])}。"
            f"来源工作表：{record['source_sheet']}；原始行：{record['source_row']}；"
            f"数据版本：{record['dataset_version_id']}。"), application_id=key)
        add(f"abstract-{key}", "paragraph", text=record["abstract"], application_id=key)
    return {"title": f"{snapshot['project_name']} 专利样本统计简报", "blocks": blocks,
            "layout": deepcopy(DEFAULT_REPORT_LAYOUT),
            "source": {key: deepcopy(value) for key, value in snapshot.items()
                       if key not in {"charts", "evidence", "representatives", "project_name"}},
            "evidence_snapshot": deepcopy(snapshot["evidence"]),
            "representative_snapshot": deepcopy(snapshot["representatives"])}


def generate_report(ctx):
    snapshot = ctx.step("report_snapshot", lambda: _load_snapshot(ctx))
    content = build_report_content(snapshot)
    if ctx.payload.get("report_mode") == "ai":
        from app.model_provider import chat

        def synthesize():
            response = chat([
                {"role": "system", "content": "请用中文总结给定专利样本事实。只引用给定编号，不新增计数，不判断法律状态或未检索的空白。"},
                {"role": "user", "content": json.dumps({"summary": snapshot["summary"],
                    "records": snapshot["representatives"]}, ensure_ascii=False)},
            ])
            return {"text": _response_text(response)}

        synthesis = ctx.step("report_synthesis", synthesize)
        content["blocks"].insert(1, {"id": str(uuid5(NAMESPACE_URL, f"{snapshot['analysis_id']}:ai-summary")),
                                    "type": "paragraph", "text": synthesis["text"], "origin": "model"})
    ctx.check_cancelled()
    with session_scope() as session:
        ctx.guard(session)
        _project(session, ctx.project_id)
        previous = session.scalar(select(ReportVersion).where(ReportVersion.run_id == ctx.run_id))
        if previous:
            return _version_result(session.get(Report, previous.report_id), previous)
        report = session.scalar(select(Report).where(Report.project_id == ctx.project_id).with_for_update())
        if not report:
            report = Report(project_id=ctx.project_id)
            session.add(report)
            session.flush()
        report.version_counter += 1
        version = ReportVersion(report_id=report.id, number=report.version_counter,
            parent_id=report.current_version_id, run_id=ctx.run_id, analysis_id=snapshot["analysis_id"],
            scope_version_id=snapshot["scope_version_id"], dataset_version_ids=snapshot["dataset_version_ids"],
            content=content, chart_ids=[chart["id"] for chart in snapshot["charts"]],
            summary="根据固定数据与范围生成统计简报", instruction="生成统计简报",
            conversation_id=getattr(ctx, "conversation_id", None))
        session.add(version)
        session.flush()
        report.current_version_id = version.id
        ctx.guard(session)
        return _version_result(report, version)


def _response_text(response):
    content = response.get("content")
    if not isinstance(content, str) or not content.strip() or len(content) > 40000:
        raise AppError(502, "INVALID_MODEL_RESPONSE", "模型没有返回有效的报告文字")
    return content.strip()


def replace_block(content, block_id, replacement):
    content = deepcopy(content)
    block = next((block for block in content["blocks"] if block["id"] == block_id), None)
    if not block:
        raise AppError(404, "BLOCK_NOT_FOUND", "指定正文块不存在")
    if block["type"] not in {"paragraph", "heading"}:
        raise AppError(422, "BLOCK_TYPE_UNSUPPORTED", "首轮仅支持替换段落或标题文字")
    block["text"] = replacement
    return content


def revise_report(ctx):
    payload = ctx.payload

    def prepare():
        with session_scope() as session:
            report = session.get(Report, payload["report_id"])
            base = session.get(ReportVersion, payload["base_version_id"])
            if not report or report.project_id != ctx.project_id or not base or base.report_id != report.id:
                raise AppError(404, "REPORT_VERSION_NOT_FOUND", "报告版本不存在")
            if report.current_version_id != payload["expected_current_version_id"]:
                raise AppError(409, "REPORT_VERSION_CONFLICT", "当前采用的版本已变化，请重新提交修改")
            block = next((row for row in base.content["blocks"] if row["id"] == payload["block_id"]), None)
            if not block:
                raise AppError(404, "BLOCK_NOT_FOUND", "指定正文块不存在")
            if block["type"] not in {"paragraph", "heading"}:
                raise AppError(422, "BLOCK_TYPE_UNSUPPORTED", "首轮仅支持替换段落或标题文字")
            return {"content": deepcopy(base.content), "analysis_id": base.analysis_id,
                    "scope_version_id": base.scope_version_id, "dataset_version_ids": base.dataset_version_ids,
                    "chart_ids": base.chart_ids, "block": deepcopy(block)}

    base = ctx.step("revision_snapshot", prepare)
    if payload.get("replacement_text") is not None:
        replacement = payload["replacement_text"]
    else:
        from app.model_provider import chat

        def rewrite():
            response = chat([
                {"role": "system", "content": "只返回修改后的单个正文块文字。遵循指令，保留证据与数据边界，不虚构事实。"},
                {"role": "user", "content": json.dumps({"instruction": payload["instruction"],
                    "block": base["block"], "source": base["content"]["source"]}, ensure_ascii=False)},
            ])
            return {"text": _response_text(response)}

        replacement = ctx.step("revision_text", rewrite)["text"]
    content = replace_block(base["content"], payload["block_id"], replacement)
    ctx.check_cancelled()
    with session_scope() as session:
        ctx.guard(session)
        _project(session, ctx.project_id)
        report = session.scalar(select(Report).where(Report.id == payload["report_id"]).with_for_update())
        previous = session.scalar(select(ReportVersion).where(ReportVersion.run_id == ctx.run_id))
        if previous:
            return _version_result(report, previous)
        if report.current_version_id != payload["expected_current_version_id"]:
            raise AppError(409, "REPORT_VERSION_CONFLICT", "修改期间当前版本发生变化，本次未覆盖其他修改")
        report.version_counter += 1
        version = ReportVersion(report_id=report.id, number=report.version_counter,
            parent_id=payload["base_version_id"], run_id=ctx.run_id, analysis_id=base["analysis_id"],
            scope_version_id=base["scope_version_id"], dataset_version_ids=base["dataset_version_ids"],
            content=content, chart_ids=base["chart_ids"], instruction=payload.get("instruction") or "替换指定文字块",
            summary=f"修改正文块 {payload['block_id']}", conversation_id=getattr(ctx, "conversation_id", None))
        session.add(version)
        session.flush()
        report.current_version_id = version.id
        ctx.guard(session)
        return _version_result(report, version)


def render_chart(option, output_path):
    renderer = Path(__file__).resolve().parents[1] / "chart-renderer" / "render.mjs"
    node = shutil.which("node")
    if not node or not renderer.is_file():
        raise AppError(503, "CHART_RENDERER_UNAVAILABLE", "图表导出环境尚未就绪")
    source = output_path.with_suffix(".json")
    source.write_text(json.dumps(option, ensure_ascii=False), encoding="utf-8")
    try:
        subprocess.run([node, str(renderer), "--input", str(source), "--output", str(output_path)],
                       check=True, capture_output=True, timeout=60)
    except (subprocess.SubprocessError, OSError) as error:
        raise AppError(502, "CHART_RENDER_FAILED", "图表渲染失败，请检查导出服务日志") from error
    if not output_path.is_file():
        raise AppError(502, "CHART_RENDER_FAILED", "图表渲染未生成图片")


def _font(run, size, family):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), family)


def build_docx(content, version_number, *, check_cancelled=lambda: None):
    """Render solely from a report version's deep snapshot, including chart options."""
    layout = {**deepcopy(DEFAULT_REPORT_LAYOUT), **content.get("layout", {})}
    margins = {**DEFAULT_REPORT_LAYOUT["margins_inches"], **layout["margins_inches"]}
    family, body_size = layout["font_family"], layout["body_font_pt"]
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(layout["page_width_inches"]), Inches(layout["page_height_inches"])
    for edge, inches in margins.items():
        setattr(section, f"{edge}_margin", Inches(inches))
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = document.styles[style_name]
        style.font.name = family
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), family)
    normal = document.styles["Normal"]
    normal.font.size = Pt(body_size)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.25
    document.add_paragraph(content["title"], "Title")
    document.add_paragraph(f"报告版本 {version_number}    分析结果 {content['source']['analysis_id']}")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    with TemporaryDirectory(prefix="patentcompass-export-") as temporary:
        for index, block in enumerate(content["blocks"]):
            check_cancelled()
            kind = block["type"]
            if kind == "heading":
                document.add_heading(block["text"], level=block.get("level", 1))
            elif kind == "paragraph":
                document.add_paragraph(block["text"])
            elif kind == "table":
                caption = document.add_paragraph(block.get("caption", ""), "Caption")
                caption.paragraph_format.keep_with_next = True
                table = document.add_table(rows=1, cols=len(block["columns"]))
                table.style = "Table Grid"
                header = table.rows[0]
                repeat = OxmlElement("w:tblHeader")
                header._tr.get_or_add_trPr().append(repeat)
                for cell, text in zip(header.cells, block["columns"]):
                    cell.text = str(text)
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8EEF5")
                    cell._tc.get_or_add_tcPr().append(shade)
                for values in block["rows"]:
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = str(value)
                for row in table.rows:
                    for cell_index, cell in enumerate(row.cells):
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell.width = Inches(5.2 if cell_index == 0 else 1.2)
                        properties = cell._tc.get_or_add_tcPr()
                        margins = OxmlElement("w:tcMar")
                        for edge in ("top", "left", "bottom", "right"):
                            margin = OxmlElement(f"w:{edge}")
                            margin.set(qn("w:w"), "90")
                            margin.set(qn("w:type"), "dxa")
                            margins.append(margin)
                        properties.append(margins)
                        borders = OxmlElement("w:tcBorders")
                        for edge in ("top", "left", "bottom", "right"):
                            border = OxmlElement(f"w:{edge}")
                            for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                                border.set(qn(f"w:{attr}"), value)
                            borders.append(border)
                        properties.append(borders)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(3)
                            paragraph.paragraph_format.space_before = Pt(3)
                            if cell_index:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif kind == "chart":
                path = Path(temporary) / f"chart-{index}.png"
                render_chart(block["chart"]["option"], path)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run().add_picture(str(path), width=Inches(layout["chart_width_inches"]))
                document.add_paragraph(block.get("caption", ""), "Caption")
                document.add_paragraph("来源：当前报告固定分析样本；计数单位为申请记录。", "Caption")
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            _font(run, 20 if paragraph.style.name == "Title" else
                  14 if paragraph.style.name == "Heading 1" else
                  11 if paragraph.style.name == "Heading 2" else body_size, family)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _font(run, body_size - 1, family)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_report(ctx):
    payload = ctx.payload
    if payload.get("format", "docx") != "docx":
        raise AppError(422, "EXPORT_FORMAT_UNSUPPORTED", "首轮仅支持 Word 导出")

    def prepare():
        with session_scope() as session:
            report = session.get(Report, payload["report_id"])
            version = session.get(ReportVersion, payload["report_version_id"])
            if not report or report.project_id != ctx.project_id or not version or version.report_id != report.id:
                raise AppError(404, "REPORT_VERSION_NOT_FOUND", "报告版本不存在")
            return {"content": deepcopy(version.content), "number": version.number,
                    "report_id": report.id, "report_version_id": version.id}

    snapshot = ctx.step("export_snapshot", prepare)
    with session_scope() as session:
        existing = session.scalar(select(StoredFile).where(StoredFile.run_id == ctx.run_id))
        if existing:
            return {"file_id": existing.id, "download_url": f"/api/files/{existing.id}",
                    "report_version_id": snapshot["report_version_id"]}
    document = build_docx(snapshot["content"], snapshot["number"], check_cancelled=ctx.check_cancelled)
    ctx.check_cancelled()
    with session_scope() as session:
        ctx.guard(session)
        file = store_bytes(session, document, f"patentcompass-report-v{snapshot['number']}.docx",
                           kind="export", project_id=ctx.project_id, run_id=ctx.run_id)
        ctx.guard(session)
        return {"file_id": file.id, "download_url": f"/api/files/{file.id}",
                "report_version_id": snapshot["report_version_id"]}
