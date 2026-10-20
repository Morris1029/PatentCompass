from datetime import timedelta
from io import BytesIO

import openpyxl
import pytest
from docx import Document
from sqlalchemy import func, select

from app import materials_service, merge_service
from app.config import get_settings
from app.data_service import import_dataset
from app.errors import AppError
from app.files import store_bytes
from app.jobs import RunContext, WaitingForInput
from app.models import (Dataset, DatasetSource, DatasetVersion, ImportConflict, Material, MaterialBlock,
                        PatentApplication, PatentPublication, ResearchRun, new_id, utcnow)


def workbook_bytes(headers, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "补充来源"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def create_base(db_session, context_factory, project):
    content = workbook_bytes(["申请号", "公开（公告）号", "申请日", "公开（公告）日", "发明名称", "摘要", "申请（专利权）人"],
                             [["WO2020CN1", "WO2021000001A1", "2020.01.01", "2021.01.01", "UAV control", "Control abstract", "ALPHA"],
                              ["WO2020CN2", "WO2021000002A1", "2020.02.01", "2021.02.01", "UAV landing", "Landing abstract", "BETA"]])
    source = store_bytes(db_session, content, "base.xlsx", "upload")
    dataset = Dataset(owner_id=get_settings().dev_owner_id, name="base", source_sha256=source.sha256)
    db_session.add(dataset)
    db_session.flush()
    base = DatasetVersion(dataset_id=dataset.id, source_file_id=source.id, source_metadata={"coverage": "partial"})
    db_session.add(base)
    db_session.flush()
    import_dataset(context_factory(kind="import", payload={"dataset_version_id": base.id}))
    db_session.expire_all()
    project.dataset_version_id = base.id
    db_session.flush()
    return base


def create_merge(db_session, base, content):
    source = store_bytes(db_session, content, "supplement.xlsx", "upload")
    version = DatasetVersion(dataset_id=base.dataset_id, source_file_id=source.id, base_version_id=base.id,
                             merge_fingerprint=source.sha256, source_metadata=dict(base.source_metadata))
    db_session.add(version)
    db_session.flush()
    return version


def test_docx_paragraph_and_table_locations_and_stable_retry(db_session, context_factory, project, tmp_path):
    document = Document()
    document.add_paragraph("无人机研究需求")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "目标"
    table.cell(0, 1).text = "自主降落"
    document.add_paragraph("下一阶段完成原型")
    target = BytesIO()
    document.save(target)
    source = store_bytes(db_session, target.getvalue(), "需求.docx", "material", project.id)
    material = Material(project_id=project.id, file_id=source.id)
    db_session.add(material)
    db_session.flush()
    ctx = context_factory(project, kind="parse_material", payload={"material_id": material.id})
    result = materials_service.parse_material(ctx)
    assert result["block_count"] == 4
    db_session.expire_all()
    first = materials_service.read_materials(db_session, [material.id], project.id)[0]
    assert [block["text"] for block in first["blocks"]] == ["无人机研究需求", "目标", "自主降落", "下一阶段完成原型"]
    assert first["blocks"][2]["location"] == {"table": 1, "row": 1, "cell": 2}
    assert all("page" not in block["location"] for block in first["blocks"])
    materials_service.parse_material(ctx)
    db_session.expire_all()
    second = materials_service.read_materials(db_session, [material.id], project.id)[0]
    assert [block["id"] for block in first["blocks"]] == [block["id"] for block in second["blocks"]]
    assert db_session.scalar(select(func.count()).select_from(MaterialBlock).where(MaterialBlock.material_id == material.id)) == 4
    with pytest.raises(AppError) as error:
        materials_service.read_materials(db_session, [material.id], new_id())
    assert error.value.code == "MATERIAL_NOT_FOUND"


def test_invalid_word_persists_failure(db_session, context_factory, project):
    source = store_bytes(db_session, b"not a docx", "broken.docx", "material", project.id)
    material = Material(project_id=project.id, file_id=source.id)
    db_session.add(material)
    db_session.flush()
    ctx = context_factory(project, kind="parse_material", payload={"material_id": material.id})
    with pytest.raises(AppError) as error:
        materials_service.parse_material(ctx)
    assert error.value.code == "INVALID_WORD_DOCUMENT"
    db_session.expire_all()
    assert material.status == "failed" and material.error["code"] == "INVALID_WORD_DOCUMENT"


def test_doc_requires_real_converter_and_preserves_original(tmp_path, monkeypatch):
    source = tmp_path / "source.doc"
    source.write_bytes(b"original binary Word")
    monkeypatch.setattr(materials_service.shutil, "which", lambda name: None)
    with pytest.raises(AppError) as error:
        materials_service._parse_source(source)
    assert error.value.code == "WORD_CONVERTER_UNAVAILABLE"
    assert source.read_bytes() == b"original binary Word"


def test_partial_enrichment_clones_old_snapshot_without_recount(db_session, context_factory, project):
    base = create_base(db_session, context_factory, project)
    content = workbook_bytes(["申请号", "IPC分类号"], [[" wo2020 cn1 ", "B64U10/13;G08G5/00"]])
    version = create_merge(db_session, base, content)
    ctx = context_factory(project, kind="merge_import", payload={"dataset_version_id": version.id})
    result = merge_service.merge_dataset(ctx)
    assert result["status"] == "ready"
    assert (result["application_count"], result["publication_count"], result["abstract_count"]) == (2, 2, 2)
    db_session.expire_all()
    merged = db_session.scalar(select(PatentApplication).where(PatentApplication.dataset_version_id == version.id,
                                                               PatentApplication.application_number == "WO2020CN1"))
    original = db_session.scalar(select(PatentApplication).where(PatentApplication.dataset_version_id == base.id,
                                                                 PatentApplication.application_number == "WO2020CN1"))
    assert merged.ipc_codes == ["B64U10/13", "G08G5/00"] and original.ipc_codes == []
    assert len(merged.raw_record["_sources"]) == 2
    assert merged.raw_record["_sources"][1]["raw_record"]["申请号"] == " wo2020 cn1 "
    assert len(db_session.scalars(select(DatasetSource).where(DatasetSource.version_id == version.id)).all()) == 2
    assert merge_service.merge_dataset(ctx)["application_count"] == 2
    assert db_session.scalar(select(func.count()).select_from(PatentApplication).where(PatentApplication.dataset_version_id == version.id)) == 2


def test_conflicts_wait_and_resume_exact_snapshot_with_explicit_choices(db_session, context_factory, project):
    base = create_base(db_session, context_factory, project)
    content = workbook_bytes(["申请号", "发明名称", "申请（专利权）人", "公开（公告）号", "公开（公告）日"],
                             [["WO2020CN1", "New control title", "GAMMA", "WO2021000001A1", "2021.01.02"]])
    version = create_merge(db_session, base, content)
    ctx = context_factory(project, kind="merge_import", payload={"dataset_version_id": version.id})
    with pytest.raises(WaitingForInput):
        merge_service.merge_dataset(ctx)
    db_session.expire_all()
    run = db_session.get(ResearchRun, ctx.run_id)
    assert run.status == "waiting_input" and run.input_request["kind"] == "import_conflicts"
    assert version.status == "waiting_input" and version.quality["unresolved_conflict_count"] == 3
    conflicts = db_session.scalars(select(ImportConflict).where(ImportConflict.version_id == version.id)).all()
    with pytest.raises(AppError):
        merge_service.resolve_conflicts(db_session, version.id, [{"conflict_id": conflicts[0].id, "choice": "new"}])
    assert all(item.resolution is None for item in conflicts)
    choices = [{"conflict_id": item.id, "choice": "merge" if item.field == "applicants" else "new"} for item in conflicts]
    merge_service.resolve_conflicts(db_session, version.id, choices)
    token = new_id()
    run.status, run.input_request, run.lease_token, run.lease_expires_at = "running", None, token, utcnow() + timedelta(hours=1)
    db_session.flush()
    result = merge_service.merge_dataset(RunContext(run.id, token))
    assert result["status"] == "ready" and result["unresolved_conflict_count"] == 0
    db_session.expire_all()
    merged = db_session.scalar(select(PatentApplication).where(PatentApplication.dataset_version_id == version.id,
                                                               PatentApplication.application_number == "WO2020CN1"))
    assert merged.title == "New control title" and merged.applicants == ["ALPHA", "GAMMA"]
    public = db_session.scalar(select(PatentPublication).where(PatentPublication.dataset_version_id == version.id,
                                                               PatentPublication.publication_number == "WO2021000001A1"))
    assert public.publication_date.isoformat() == "2021-01-02"
    original = db_session.scalar(select(PatentApplication).where(PatentApplication.dataset_version_id == base.id,
                                                                 PatentApplication.application_number == "WO2020CN1"))
    assert original.title == "UAV control" and original.applicants == ["ALPHA"]


def test_new_partial_application_fails_whole_merge(db_session, context_factory, project):
    base = create_base(db_session, context_factory, project)
    version = create_merge(db_session, base, workbook_bytes(["申请号", "IPC分类号"], [["WO2099CN9", "B64U"]]))
    ctx = context_factory(project, kind="merge_import", payload={"dataset_version_id": version.id})
    with pytest.raises(AppError) as error:
        merge_service.merge_dataset(ctx)
    assert error.value.code == "DATASET_QUALITY_FAILED"
    db_session.expire_all()
    assert version.status == "failed" and version.quality["errors"]
    assert db_session.scalar(select(func.count()).select_from(PatentApplication).where(PatentApplication.dataset_version_id == version.id)) == 0


def test_publication_cannot_move_to_another_application(db_session, context_factory, project):
    base = create_base(db_session, context_factory, project)
    version = create_merge(db_session, base, workbook_bytes(["申请号", "公开（公告）号", "公开（公告）日"],
                                                           [["WO2020CN2", "WO2021000001A1", "2021.01.01"]]))
    with pytest.raises(AppError) as error:
        merge_service.merge_dataset(context_factory(project, kind="merge_import", payload={"dataset_version_id": version.id}))
    assert error.value.code == "PUBLICATION_APPLICATION_CONFLICT"
    assert db_session.scalar(select(func.count()).select_from(PatentApplication).where(PatentApplication.dataset_version_id == version.id)) == 0
