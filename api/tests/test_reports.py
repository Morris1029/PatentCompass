from copy import deepcopy
from datetime import date
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from sqlalchemy import func, select

from app.errors import AppError
from app.config import get_settings
from app.files import resolve_file
from app.models import (
    AnalysisResult, Chart, Dataset, DatasetVersion, Evidence, PatentApplication,
    PatentPublication, Report, ReportVersion, StoredFile,
)
from app.report_service import build_docx, export_report, generate_report, replace_block, revise_report


@pytest.fixture
def report_analysis(db_session, project, context_factory):
    source = StoredFile(owner_id=project.owner_id, kind="upload", key="upload/report-test.xlsx",
                        original_name="report-test.xlsx", sha256="a" * 64, size=1,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    dataset = Dataset(owner_id=project.owner_id, name="报告测试", source_sha256="a" * 64)
    db_session.add_all([source, dataset])
    db_session.flush()
    version = DatasetVersion(dataset_id=dataset.id, source_file_id=source.id, status="ready",
                             quality={"application_count": 1, "publication_count": 1, "abstract_count": 1},
                             source_metadata={"query": "无人机", "region": "中国", "partial_results": True})
    db_session.add(version)
    db_session.flush()
    project.dataset_version_id = version.id
    patent = PatentApplication(dataset_version_id=version.id, application_number="PCT/TEST/1",
        application_date=date(2024, 1, 1), title="Unmanned aerial vehicle",
        abstract="An unmanned aerial vehicle uses a visual sensor for landing.",
        applicants=["Applicant A", "Applicant B"], ipc_codes=["B64U10/00", "G05D1/00"],
        source_sheet="Patents", source_row=2)
    db_session.add(patent)
    db_session.flush()
    publication = PatentPublication(application_id=patent.id, dataset_version_id=version.id,
        publication_number="WO2024000001A1", publication_date=patent.application_date, kind_code="A1")
    db_session.add(publication)
    db_session.flush()
    context = context_factory(project, "analyze", {})
    analysis = AnalysisResult(project_id=project.id, run_id=context.run_id,
        scope_version_id=project.scope_version_id, dataset_version_ids=[version.id], patent_ids=[patent.id],
        summary={"application_count": 1, "publication_count": 1, "abstract_count": 1,
                 "yearly": [{"year": 2024, "count": 1}],
                 "applicants": [{"name": "Applicant A", "count": 1}, {"name": "Applicant B", "count": 1}],
                 "ipc": [{"code": "B64U10/00", "count": 1}, {"code": "G05D1/00", "count": 1}]}, views={})
    db_session.add(analysis)
    db_session.flush()
    project.current_analysis_id = analysis.id
    for key in ("yearly", "applicants", "ipc"):
        db_session.add(Chart(project_id=project.id, analysis_id=analysis.id, kind=key,
            option={"xAxis": {"type": "category", "data": ["2024"]}, "yAxis": {"type": "value"},
                    "series": [{"type": "bar", "data": [1]}]},
            data={"title": key, "rows": [{"label": "2024", "value": "2024", "count": 1}]},
            provenance={"dataset_version_ids": [version.id], "application_ids": [patent.id]}))
    db_session.add(Evidence(project_id=project.id, analysis_id=analysis.id, application_id=patent.id,
                            kind="patent", content={"title": patent.title, "abstract": patent.abstract,
                                                    "source_sheet": "Patents", "source_row": 2}))
    db_session.flush()
    return analysis


def _text_block(version):
    return next(block for block in version.content["blocks"] if block["type"] == "paragraph")


def test_report_versions_branch_without_mutating_prior_content(db_session, project, context_factory, report_analysis):
    context = context_factory(project, "generate_report", {"analysis_id": report_analysis.id})
    first = generate_report(context)
    assert generate_report(context)["report_version_id"] == first["report_version_id"]
    v1 = db_session.get(ReportVersion, first["report_version_id"])
    frozen = deepcopy(v1.content)
    block_id = _text_block(v1)["id"]

    def revise(base, expected, text):
        return revise_report(context_factory(project, "revise_report", {
            "report_id": first["report_id"], "base_version_id": base,
            "expected_current_version_id": expected, "block_id": block_id, "replacement_text": text}))

    second = revise(v1.id, v1.id, "Second version")
    db_session.expire_all()
    report = db_session.get(Report, first["report_id"])
    report.current_version_id = v1.id
    db_session.flush()
    third = revise(v1.id, v1.id, "Third version branched from first")
    db_session.expire_all()
    v2 = db_session.get(ReportVersion, second["report_version_id"])
    v3 = db_session.get(ReportVersion, third["report_version_id"])
    assert (v1.number, v2.number, v3.number) == (1, 2, 3)
    assert v2.parent_id == v3.parent_id == v1.id
    assert db_session.get(ReportVersion, v1.id).content == frozen
    assert _text_block(v2)["text"] == "Second version"
    assert _text_block(v3)["text"] == "Third version branched from first"
    assert [b for b in v3.content["blocks"] if b["type"] == "chart"] == [
        b for b in frozen["blocks"] if b["type"] == "chart"]
    assert v3.content["evidence_snapshot"] == frozen["evidence_snapshot"]


def test_revision_detects_adoption_change_before_writing(db_session, project, context_factory, report_analysis):
    first = generate_report(context_factory(project, "generate_report", {"analysis_id": report_analysis.id}))
    original = db_session.get(ReportVersion, first["report_version_id"])
    block_id = _text_block(original)["id"]
    second = revise_report(context_factory(project, "revise_report", {
        "report_id": first["report_id"], "base_version_id": original.id,
        "expected_current_version_id": original.id, "block_id": block_id, "replacement_text": "V2"}))
    context = context_factory(project, "revise_report", {
        "report_id": first["report_id"], "base_version_id": second["report_version_id"],
        "expected_current_version_id": second["report_version_id"], "block_id": block_id,
        "replacement_text": "This must not be persisted"})
    actual_step = context.step

    def step(key, callback):
        result = actual_step(key, callback)
        if key == "revision_snapshot":
            db_session.expire_all()
            db_session.get(Report, first["report_id"]).current_version_id = original.id
            db_session.flush()
        return result

    context.step = step
    with pytest.raises(AppError) as failure:
        revise_report(context)
    assert failure.value.status_code == 409
    assert failure.value.code == "REPORT_VERSION_CONFLICT"
    db_session.expire_all()
    assert db_session.get(Report, first["report_id"]).version_counter == 2
    assert db_session.scalar(select(func.count()).select_from(ReportVersion).where(
        ReportVersion.report_id == first["report_id"])) == 2


def test_report_freezes_source_and_chart_data(db_session, project, context_factory, report_analysis):
    first = generate_report(context_factory(project, "generate_report", {"analysis_id": report_analysis.id}))
    frozen = deepcopy(db_session.get(ReportVersion, first["report_version_id"]).content)
    chart = db_session.scalar(select(Chart).where(Chart.analysis_id == report_analysis.id))
    chart.option = {"series": [{"type": "bar", "data": [999]}]}
    report_analysis.summary = {"application_count": 999}
    db_session.flush()
    db_session.expire_all()
    assert db_session.get(ReportVersion, first["report_version_id"]).content == frozen
    assert frozen["source"]["summary"]["application_count"] == 1
    assert all(block["chart"]["option"]["series"][0]["data"] == [1]
               for block in frozen["blocks"] if block["type"] == "chart")


def test_replace_block_preserves_identity_and_rejects_chart_edits():
    content = {"blocks": [{"id": "p", "type": "paragraph", "text": "Before"},
                          {"id": "c", "type": "chart", "chart": {"option": {}}}]}
    result = replace_block(content, "p", "After")
    assert content["blocks"][0]["text"] == "Before"
    assert result["blocks"][0] == {"id": "p", "type": "paragraph", "text": "After"}
    with pytest.raises(AppError) as failure:
        replace_block(content, "c", "Change chart")
    assert failure.value.status_code == 422


def test_natural_language_revision_without_model_does_not_create_version(db_session, project, context_factory,
                                                                      report_analysis, monkeypatch):
    monkeypatch.setattr(get_settings(), "llm_base_url", "")
    first = generate_report(context_factory(project, "generate_report", {"analysis_id": report_analysis.id}))
    version = db_session.get(ReportVersion, first["report_version_id"])
    context = context_factory(project, "revise_report", {
        "report_id": first["report_id"], "base_version_id": version.id,
        "expected_current_version_id": version.id, "block_id": _text_block(version)["id"],
        "instruction": "请简化这段话"})
    with pytest.raises(AppError) as failure:
        revise_report(context)
    assert failure.value.code == "MODEL_NOT_CONFIGURED"
    assert failure.value.status_code == 503
    db_session.expire_all()
    assert db_session.get(Report, first["report_id"]).version_counter == 1


def test_export_uses_requested_version_and_reuses_existing_file(db_session, project, context_factory,
                                                             report_analysis, monkeypatch):
    first = generate_report(context_factory(project, "generate_report", {"analysis_id": report_analysis.id}))
    version = db_session.get(ReportVersion, first["report_version_id"])
    frozen = deepcopy(version.content)
    revise_report(context_factory(project, "revise_report", {
        "report_id": first["report_id"], "base_version_id": version.id,
        "expected_current_version_id": version.id, "block_id": _text_block(version)["id"],
        "replacement_text": "This is the new current version, which must not enter the v1 export"}))
    calls = []

    def build(content, number, **kwargs):
        calls.append((deepcopy(content), number))
        output = BytesIO()
        Document().save(output)
        return output.getvalue()

    monkeypatch.setattr("app.report_service.build_docx", build)
    context = context_factory(project, "export", {
        "report_id": first["report_id"], "report_version_id": version.id, "format": "docx"})
    result = export_report(context)
    assert export_report(context) == result
    assert result["report_version_id"] == version.id
    assert calls == [(frozen, 1)]
    file = db_session.get(StoredFile, result["file_id"])
    assert file.run_id == context.run_id
    assert resolve_file(file.key).read_bytes().startswith(b"PK")


def test_docx_contains_frozen_text_tables_and_three_chart_images(db_session, project, context_factory,
                                                               report_analysis, monkeypatch):
    first = generate_report(context_factory(project, "generate_report", {"analysis_id": report_analysis.id}))
    version = db_session.get(ReportVersion, first["report_version_id"])
    seen_options = []

    def render(option, path):
        # A valid tiny PNG isolates Word packaging from the separately verified ECharts renderer.
        import base64
        seen_options.append(deepcopy(option))
        path.write_bytes(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="))

    monkeypatch.setattr("app.report_service.render_chart", render)
    result = build_docx(version.content, version.number)
    document = Document(BytesIO(result))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "专利样本统计简报" in text
    assert "1 条申请记录" in text
    assert "WO2024000001A1" in text
    assert "visual sensor for landing" in text
    assert len(document.tables) == 3
    assert len(document.inline_shapes) == len(seen_options) == 3
    with ZipFile(BytesIO(result)) as archive:
        assert "word/media/image1.png" in archive.namelist()
        assert b"PAGE" in archive.read("word/footer1.xml")


def test_revision_and_export_preserve_frozen_document_layout(db_session, project, context_factory,
                                                           report_analysis, monkeypatch):
    from app import report_service
    first = generate_report(context_factory(project, "generate_report", {"analysis_id": report_analysis.id}))
    v1 = db_session.get(ReportVersion, first["report_version_id"])
    frozen_layout = deepcopy(v1.content["layout"])
    second = revise_report(context_factory(project, "revise_report", {
        "report_id": first["report_id"], "base_version_id": v1.id,
        "expected_current_version_id": v1.id, "block_id": _text_block(v1)["id"],
        "replacement_text": "Layout must remain unchanged"}))
    v2 = db_session.get(ReportVersion, second["report_version_id"])
    assert v2.content["layout"] == frozen_layout
    # A later template default must not change an already saved version's pages or font.
    monkeypatch.setattr(report_service, "DEFAULT_REPORT_LAYOUT", {
        **frozen_layout, "page_width_inches": 7.0, "body_font_pt": 14, "font_family": "Arial"})
    content = deepcopy(v2.content)
    content["blocks"] = [block for block in content["blocks"] if block["type"] != "chart"]
    document = Document(BytesIO(build_docx(content, v2.number)))
    assert round(document.sections[0].page_width.inches, 2) == frozen_layout["page_width_inches"]
    assert round(document.sections[0].page_height.inches, 2) == frozen_layout["page_height_inches"]
    assert round(document.sections[0].left_margin.inches, 2) == frozen_layout["margins_inches"]["left"]
    assert document.styles["Normal"].font.name == frozen_layout["font_family"]
    assert document.styles["Normal"].font.size.pt == frozen_layout["body_font_pt"]
    del content["layout"]
    legacy = Document(BytesIO(build_docx(content, v2.number)))
    assert legacy.sections[0].page_width.inches == 7.0
    assert legacy.styles["Normal"].font.size.pt == 14
