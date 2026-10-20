import hashlib
from pathlib import Path

import httpx
import openpyxl
import pytest

from app import model_provider
from app import data_service
from app.config import Settings
from app.data_service import HEADERS, _keyword_groups, _priorities, parse_workbook
from app.errors import AppError


SAMPLE = Path(__file__).resolve().parents[2] / "database" / "patent-demo.xlsx"


def write_workbook(path, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def row():
    return ["WO2020CN12345", "WO2021000123A1;WO2021000123A9", "2020.01.02", "2021.01.07;2021.02.04",
            "B64U10/13;G08G5/00", "ALPHA LTD; BETA LTD", "PERSON A; PERSON B", "UAV landing control",
            "CN201910001234A 20190101", "2019.01.01", "", "CN", "An autonomous landing method."]


def test_real_sample_counts_and_source_preserved():
    before = hashlib.sha256(SAMPLE.read_bytes()).hexdigest()
    parsed = parse_workbook(SAMPLE)
    quality = parsed["quality"]
    assert quality["status"] == "ready"
    assert (quality["application_count"], quality["publication_count"], quality["abstract_count"]) == (520, 526, 520)
    assert quality["source_row_count"] == 520
    assert quality["missing_fields"] == {"document_type": 520}
    assert quality["publication_prefixes"] == {"WO": 526}
    assert quality["errors"] == []
    assert sum(len(record["publications"]) == 2 for record in parsed["applications"]) == 6
    assert all("CN" in record["countries"] for record in parsed["applications"])
    assert any(len(record["countries"]) > 1 for record in parsed["applications"])
    assert hashlib.sha256(SAMPLE.read_bytes()).hexdigest() == before


def test_multiple_publications_and_priorities_are_preserved(tmp_path):
    path = tmp_path / "source.xlsx"
    source = row()
    source[8] = "CN202010000001A 20200101\n\nCN202010000002A 20200101\n\nCN202010000003A 20200103"
    source[9] = "2020.01.01;2020.01.03"
    write_workbook(path, [source])
    parsed = parse_workbook(path)
    assert parsed["quality"]["status"] == "ready"
    patent = parsed["applications"][0]
    assert [p["publication_date"] for p in patent["publications"]] == ["2021-01-07", "2021-02-04"]
    assert [p["date"] for p in patent["priorities"]] == ["2020-01-01", "2020-01-01", "2020-01-03"]
    assert patent["raw_record"]["优先权号"] == source[8]
    assert patent["source_row"] == 2


def test_duplicate_or_invalid_rows_fail_entire_version_without_silent_drop(tmp_path):
    path = tmp_path / "duplicates.xlsx"
    invalid = row()
    invalid[0] = "WO2020CN55555"
    invalid[3] = "2021.01.07"
    write_workbook(path, [row(), row(), invalid])
    quality = parse_workbook(path)["quality"]
    assert quality["status"] == "failed"
    assert quality["source_row_count"] == 3
    assert len(quality["errors"]) == 2
    assert quality["errors"][0]["row"] == 3
    assert "重复" in quality["errors"][0]["message"]
    assert "数量不一致" in quality["errors"][1]["message"]


def test_priority_embedded_dates_are_checked():
    with pytest.raises(ValueError, match="不一致"):
        _priorities("CN202010000001A 20200101", "2020.01.02")


def test_empty_workbook_is_not_ready(tmp_path):
    path = tmp_path / "empty.xlsx"
    write_workbook(path, [])
    assert parse_workbook(path)["quality"]["status"] == "failed"


def test_chinese_keyword_expansion_keeps_all_requested_concepts():
    groups = _keyword_groups("无人机 自主降落")
    assert ["无人机", "unmanned aerial vehicle", "uav", "drone"] in groups
    assert ["自主", "autonomous", "automatic"] in groups
    assert ["降落", "landing", "land"] in groups


def test_unconfigured_models_do_not_return_simulated_success(monkeypatch):
    monkeypatch.setattr(model_provider, "get_settings", lambda: Settings(_env_file=None))
    with pytest.raises(AppError) as error:
        model_provider.chat([{"role": "user", "content": "hello"}])
    assert error.value.code == "MODEL_NOT_CONFIGURED"
    with pytest.raises(AppError):
        model_provider.embed(["unmanned aerial vehicle"])


def test_provider_embedding_order_and_message(monkeypatch):
    settings = Settings(_env_file=None, llm_base_url="https://models.invalid/v1", llm_api_key="test-key", llm_model="chat-test",
                        embedding_base_url="https://models.invalid/v1", embedding_api_key="test-key", embedding_model="embed-test")
    monkeypatch.setattr(model_provider, "get_settings", lambda: settings)

    def post(url, json, headers, timeout):
        assert headers["Authorization"] == "Bearer test-key"
        if url.endswith("embeddings"):
            assert json["input"] == ["first", "second"]
            result = {"data": [{"index": 1, "embedding": [0, 1]}, {"index": 0, "embedding": [1, 0]}]}
        else:
            result = {"choices": [{"message": {"role": "assistant", "content": "response"}}]}
        return httpx.Response(200, request=httpx.Request("POST", url), json=result)

    monkeypatch.setattr(model_provider.httpx, "post", post)
    assert model_provider.embed(["first", "second"]) == [[1, 0], [0, 1]]
    assert model_provider.chat([{"role": "user", "content": "request"}])["content"] == "response"


@pytest.mark.parametrize("vectors", [[{"index": 0, "embedding": [0, 0]}], [{"index": 0, "embedding": [float("inf"), 1]}], []])
def test_invalid_provider_vectors_are_rejected(monkeypatch, vectors):
    monkeypatch.setattr(model_provider, "_post", lambda *args: {"data": vectors})
    with pytest.raises(AppError) as error:
        model_provider.embed(["source"])
    assert error.value.code == "MODEL_RESPONSE_INVALID"


@pytest.fixture
def imported_sample(db_session, project, context_factory, monkeypatch):
    from app.models import Dataset, DatasetVersion, StoredFile
    digest = hashlib.sha256(SAMPLE.read_bytes()).hexdigest()
    source = StoredFile(owner_id=project.owner_id, kind="upload", key="upload/sample.xlsx",
                        original_name="patent-demo.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        sha256=digest, size=SAMPLE.stat().st_size)
    dataset = Dataset(owner_id=project.owner_id, name="无人机样本", source_sha256=digest)
    db_session.add_all([source, dataset])
    db_session.flush()
    version = DatasetVersion(dataset_id=dataset.id, source_file_id=source.id,
                             source_metadata={"query": "无人机", "region": "中国", "partial_results": True})
    db_session.add(version)
    db_session.flush()
    monkeypatch.setattr(data_service, "resolve_file", lambda key: SAMPLE)
    context = context_factory(project, "import", {"dataset_version_id": version.id})
    result = data_service.import_dataset(context)
    assert result["status"] == "ready"
    assert data_service.import_dataset(context)["dataset_version_id"] == version.id
    db_session.expire_all()
    project.dataset_version_id = version.id
    db_session.flush()
    return version


def test_import_statistics_match_independent_spreadsheet_counts(db_session, project, context_factory, imported_sample):
    from collections import Counter
    from sqlalchemy import func, select
    from app.models import AnalysisResult, Chart, PatentApplication, PatentPublication
    context = context_factory(project, "analyze", {})
    result = data_service.analyze(context)
    assert data_service.analyze(context) == result
    db_session.expire_all()
    analysis = db_session.get(AnalysisResult, result["analysis_id"])
    assert analysis.summary["application_count"] == 520
    assert analysis.summary["publication_count"] == 526
    assert db_session.scalar(select(func.count(PatentApplication.id))) == 520
    assert db_session.scalar(select(func.count(PatentPublication.id))) == 526
    assert project.current_analysis_id == analysis.id
    workbook = openpyxl.load_workbook(SAMPLE, read_only=True, data_only=True)
    expected_applicants, expected_ipc, expected_year, expected_countries = Counter(), Counter(), Counter(), Counter()
    for row_value in list(workbook.active.iter_rows(values_only=True))[1:]:
        expected_applicants.update(set(value.strip() for value in row_value[5].split(";")))
        expected_ipc.update(set(value.strip() for value in row_value[4].split(";")))
        expected_year.update([int(row_value[2][:4])])
        expected_countries.update(set(data_service.split_values(row_value[11])))
    workbook.close()
    assert {r["name"]: r["count"] for r in analysis.summary["applicants"]} == expected_applicants
    assert {r["code"]: r["count"] for r in analysis.summary["ipc"]} == expected_ipc
    assert {r["year"]: r["count"] for r in analysis.summary["yearly"]} == expected_year
    assert {r["country"]: r["count"] for r in analysis.summary["country"]} == expected_countries == {"CN": 520, "JP": 1, "TJ": 1}
    assert analysis.views["geography"]["status"] == "ready"
    assert set(analysis.views["geography"]["data"]["unavailable"]) == {"province", "protection", "family"}
    chart = db_session.scalar(select(Chart).where(Chart.analysis_id == analysis.id, Chart.kind == "applicants"))
    name = chart.data["rows"][0]["value"]
    historical = data_service.search_patents(db_session, project, {"application_ids": chart.selections["values"][name]})
    assert len(historical["items"]) == expected_applicants[name]
    assert all(name in item.applicants for item in historical["items"])
    assert analysis.views["opportunities"]["status"] == "unavailable"


def test_explicit_filters_and_run_scope_snapshot(db_session, project, context_factory, imported_sample):
    from app.models import AnalysisResult, ScopeVersion
    captured = context_factory(project, "analyze", {})
    original_scope = db_session.get(ScopeVersion, project.scope_version_id)
    new_scope = ScopeVersion(project_id=project.id, number=original_scope.number + 1, confirmed=True,
                             content={"application_start": "2024-01-01", "application_end": "2024-12-31"})
    db_session.add(new_scope)
    db_session.flush()
    project.scope_version_id = new_scope.id
    db_session.flush()
    result = data_service.analyze(captured)
    db_session.expire_all()
    assert db_session.get(AnalysisResult, result["analysis_id"]).summary["application_count"] == 520
    assert project.current_analysis_id is None
    filtered = data_service.search_patents(db_session, project, {"ipc": "B64U10/13"})
    assert all(item.application_date.year == 2024 and any("B64U10/13" in code for code in item.ipc_codes)
               for item in filtered["items"])


def test_semantic_search_requires_index_and_uses_pgvector(db_session, project, context_factory, imported_sample, monkeypatch):
    from sqlalchemy import select
    from app.models import PatentApplication
    settings = Settings(_env_file=None, embedding_base_url="https://test.invalid/v1", embedding_api_key="test", embedding_model="controlled-test")
    monkeypatch.setattr(data_service, "get_settings", lambda: settings)
    monkeypatch.setattr(data_service, "embed", lambda texts: [[1., 0.] if "landing" in text.lower() else [0., 1.] for text in texts])
    with pytest.raises(AppError) as error:
        data_service.search_patents(db_session, project, {"mode": "semantic", "q": "landing"})
    assert error.value.code == "EMBEDDING_INDEX_NOT_READY"
    context = context_factory(project, "index", {})
    assert data_service.build_embeddings(context)["indexed_count"] == 520
    assert data_service.build_embeddings(context)["indexed_count"] == 520
    found = data_service.search_patents(db_session, project, {"mode": "semantic", "q": "landing", "candidate_limit": 3})
    assert found["scope_application_count"] == 520
    assert found["candidate_count"] == 3
    assert all("landing" in f"{item.title} {item.abstract}".lower() for item in found["items"])
    limited_ids = list(db_session.scalars(select(PatentApplication.id).where(PatentApplication.dataset_version_id == imported_sample.id).limit(2)))
    limited = data_service.search_patents(db_session, project, {"mode": "semantic", "q": "landing", "limit": 3,
                                                              "application_ids": limited_ids})
    assert {item.id for item in limited["items"]} == set(limited_ids)


def test_embedding_batches_follow_config_and_checkpoint_model_and_boundaries(db_session, project, context_factory, monkeypatch):
    from datetime import date
    from sqlalchemy import func, select
    from app.models import Dataset, DatasetVersion, PatentApplication, PatentEmbedding, RunStep, StoredFile

    source = StoredFile(owner_id=project.owner_id, kind="upload", key="upload/batch.xlsx", original_name="batch.xlsx",
                        media_type="application/octet-stream", sha256="a" * 64, size=1)
    dataset = Dataset(owner_id=project.owner_id, name="Embedding batch sample", source_sha256="a" * 64)
    db_session.add_all([source, dataset])
    db_session.flush()
    version = DatasetVersion(dataset_id=dataset.id, source_file_id=source.id, status="ready")
    db_session.add(version)
    db_session.flush()
    db_session.add_all([PatentApplication(dataset_version_id=version.id, application_number=f"BATCH{index}",
                                         application_date=date(2020, 1, 1), title=f"UAV {index}", abstract="Landing control",
                                         source_sheet="Sheet1", source_row=index + 2) for index in range(25)])
    project.dataset_version_id = version.id
    db_session.flush()
    settings = Settings(_env_file=None, embedding_base_url="https://test.invalid/v1", embedding_api_key="test",
                        embedding_model="batch-model", embedding_batch_size=10)
    monkeypatch.setattr(data_service, "get_settings", lambda: settings)
    sizes = []

    def controlled_embed(texts):
        if texts:
            sizes.append(len(texts))
        return [[1., 0.] for _ in texts]

    monkeypatch.setattr(data_service, "embed", controlled_embed)
    context = context_factory(project, "index", {})
    assert data_service.build_embeddings(context)["indexed_count"] == 25
    assert sizes == [10, 10, 5]
    data_service.build_embeddings(context)
    assert sizes == [10, 10, 5]

    settings.embedding_batch_size = 7
    data_service.build_embeddings(context)
    assert sizes == [10, 10, 5]  # Existing vectors are reused, while new checkpoint boundaries are recorded.
    assert db_session.scalar(select(func.count(RunStep.id)).where(RunStep.run_id == context.run_id)) == 7

    settings.embedding_model = "another-model"
    settings.embedding_batch_size = 10
    data_service.build_embeddings(context)
    assert sizes == [10, 10, 5, 10, 10, 5]
    assert db_session.scalar(select(func.count(PatentEmbedding.id))) == 50
