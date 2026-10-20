from copy import deepcopy
from io import BytesIO

from docx import Document
import pytest
from sqlalchemy import func, select

from app import chart_service, report_service
from app.errors import AppError
from app.models import Chart, ChartVersion, Conversation, Evidence, Report, ReportVersion, ResearchRun, new_id
from test_reports import report_analysis


def generated(db_session, project, context_factory, analysis):
    result = report_service.generate_report(context_factory(project, "generate_report", {"analysis_id": analysis.id}))
    return db_session.get(ReportVersion, result["report_version_id"])


def revise(db_session, project, context_factory, base, operations, **extra):
    result = report_service.revise_report(context_factory(project, "revise_report", {
        "report_id": base.report_id, "base_version_id": base.id, "expected_current_version_id": base.id,
        "operations": operations, **extra}))
    db_session.expire_all()
    return db_session.get(ReportVersion, result["report_version_id"])


def test_multiple_text_title_insert_move_remove_and_layout_are_one_version(db_session, project, context_factory, report_analysis):
    base = generated(db_session, project, context_factory, report_analysis)
    frozen = deepcopy(base.content)
    texts = [block for block in base.content["blocks"] if block["type"] == "paragraph"][:2]
    chart = next(block for block in base.content["blocks"] if block["type"] == "chart")
    conversation = Conversation(project_id=project.id, title="来源对话")
    db_session.add(conversation)
    db_session.flush()
    current = revise(db_session, project, context_factory, base, [
        {"op": "set_title", "text": "无人机导航报告"},
        {"op": "replace_text", "block_id": texts[0]["id"], "text": "第一段修订"},
        {"op": "replace_text", "block_id": texts[1]["id"], "text": "第二段修订"},
        {"op": "insert_block", "block": {"id": "new-section", "type": "heading", "text": "新增章节", "level": 1}},
        {"op": "insert_block", "after_id": "new-section", "block": {"id": "new-paragraph", "type": "paragraph", "text": "新增正文"}},
        {"op": "move_block", "block_id": "new-paragraph", "after_id": texts[0]["id"]},
        {"op": "remove_block", "block_id": "new-section"},
        {"op": "set_chart_layout", "block_id": chart["id"], "layout": {"width_percent": 50, "alignment": "right"}},
        {"op": "set_layout", "layout": {"body_font_pt": 12, "margins_inches": {"left": 1}}},
    ], conversation_id=conversation.id)
    assert current.number == 2 and current.parent_id == base.id and current.conversation_id == conversation.id
    assert current.content["title"] == "无人机导航报告"
    blocks = current.content["blocks"]
    assert next(block for block in blocks if block["id"] == texts[0]["id"])["text"] == "第一段修订"
    assert next(block for block in blocks if block["id"] == texts[1]["id"])["text"] == "第二段修订"
    assert [block["id"] for block in blocks].index("new-paragraph") == [block["id"] for block in blocks].index(texts[0]["id"]) + 1
    assert "new-section" not in [block["id"] for block in blocks]
    assert next(block for block in blocks if block["id"] == chart["id"])["layout"] == {"width_percent": 50, "alignment": "right"}
    assert current.content["layout"]["body_font_pt"] == 12 and current.content["layout"]["margins_inches"]["left"] == 1
    assert current.content["layout"]["margins_inches"]["right"] == frozen["layout"]["margins_inches"]["right"]
    assert db_session.get(ReportVersion, base.id).content == frozen


def test_invalid_late_operation_cannot_partially_persist_revision(db_session, project, context_factory, report_analysis):
    base = generated(db_session, project, context_factory, report_analysis)
    frozen = deepcopy(base.content)
    with pytest.raises(AppError) as error:
        revise(db_session, project, context_factory, base, [
            {"op": "set_title", "text": "Must not survive"},
            {"op": "insert_block", "block": {"type": "paragraph", "text": "Must not survive"}},
            {"op": "remove_block", "block_id": "missing-block"},
        ])
    assert error.value.code == "BLOCK_NOT_FOUND"
    db_session.expire_all()
    assert db_session.get(ReportVersion, base.id).content == frozen
    assert db_session.get(Report, base.report_id).current_version_id == base.id
    assert db_session.scalar(select(func.count()).select_from(ReportVersion).where(ReportVersion.report_id == base.report_id)) == 1


def test_insert_update_chart_freezes_requested_history_independent_of_adoption(db_session, project, context_factory, report_analysis):
    base = generated(db_session, project, context_factory, report_analysis)
    chart = db_session.scalar(select(Chart).where(Chart.analysis_id == report_analysis.id))
    old = chart_service.snapshot(db_session, chart)
    new = ChartVersion(chart_id=chart.id, number=2, option={"series": [{"type": "bar", "data": [2]}]},
                       data={**deepcopy(chart.data), "title": "新版图表"}, data_sha256="b" * 64,
                       provenance=deepcopy(chart.provenance), selections=deepcopy(chart.selections), interpretation={})
    db_session.add(new)
    db_session.flush()
    chart_service.adopt(db_session, chart, new.id, old["version_id"])
    original = next(block for block in base.content["blocks"] if block["type"] == "chart" and block["chart_id"] == chart.id)
    updated = revise(db_session, project, context_factory, base, [
        {"op": "insert_chart", "chart_id": chart.id, "chart_version_id": old["version_id"], "after_id": original["id"]},
        {"op": "update_chart", "block_id": original["id"], "chart_id": chart.id, "chart_version_id": new.id},
    ])
    blocks = [block for block in updated.content["blocks"] if block["type"] == "chart" and block["chart_id"] == chart.id]
    assert len(blocks) == 2
    assert blocks[0]["chart_version_id"] == new.id and blocks[0]["chart"]["option"]["series"][0]["data"] == [2]
    assert blocks[1]["chart_version_id"] == old["version_id"] and blocks[1]["chart"]["option"]["series"][0]["data"] == [1]
    chart_service.adopt(db_session, chart, old["version_id"], new.id)
    db_session.expire_all()
    assert next(block for block in db_session.get(ReportVersion, updated.id).content["blocks"] if block["id"] == original["id"])["chart_version_id"] == new.id
    assert next(block for block in db_session.get(ReportVersion, base.id).content["blocks"] if block["id"] == original["id"])["chart_version_id"] == old["version_id"]


def test_chart_and_conversation_cross_project_references_are_rejected(db_session, project, context_factory, report_analysis):
    base = generated(db_session, project, context_factory, report_analysis)
    with pytest.raises(AppError) as error:
        revise(db_session, project, context_factory, base, [{"op": "insert_chart", "chart_id": new_id()}])
    assert error.value.code == "CHART_NOT_FOUND"
    with pytest.raises(AppError) as error:
        revise(db_session, project, context_factory, base, [{"op": "set_title", "text": "title"}], conversation_id=new_id())
    assert error.value.code == "CONVERSATION_NOT_FOUND"


def test_roadmap_uses_frozen_analysis_and_stable_action_ids(db_session, project, context_factory, report_analysis):
    report_analysis.views = {"roadmap": {"status": "ready", "data": {"items": [{
        "id": "direction-fixed", "title": "视觉降落核查", "priority": "P1", "conditions": ["核对原文"],
        "supporting_application_ids": report_analysis.patent_ids,
        "phases": [{"id": "phase-fixed", "title": "样本核查", "actions": [{"id": "action-fixed", "title": "比较相近申请", "owner": "待分配"}]}],
    }]}}}
    db_session.flush()
    base = generated(db_session, project, context_factory, report_analysis)
    report_analysis.views = {"roadmap": {"status": "ready", "data": {"items": []}}}
    db_session.flush()
    current = revise(db_session, project, context_factory, base, [{"op": "append_roadmap"}])
    roadmap = [block for block in current.content["blocks"] if block.get("direction_id") == "direction-fixed"]
    assert any(block.get("action_id") == "action-fixed" for block in roadmap)
    assert any("比较相近申请" in block.get("text", "") for block in roadmap)
    assert any(block.get("section") == "roadmap" for block in current.content["blocks"])
    with pytest.raises(AppError) as error:
        revise(db_session, project, context_factory, current, [{"op": "append_roadmap"}])
    assert error.value.code == "ROADMAP_ALREADY_PRESENT"


def test_unavailable_roadmap_and_invalid_layout_never_fabricate_content(db_session, project, context_factory, report_analysis):
    base = generated(db_session, project, context_factory, report_analysis)
    with pytest.raises(AppError) as error:
        revise(db_session, project, context_factory, base, [{"op": "append_roadmap"}])
    assert error.value.code == "ROADMAP_UNAVAILABLE"
    chart = next(block for block in base.content["blocks"] if block["type"] == "chart")
    with pytest.raises(AppError) as error:
        revise(db_session, project, context_factory, base, [{"op": "set_chart_layout", "block_id": chart["id"], "layout": {"width_percent": 200}}])
    assert error.value.code == "INVALID_REPORT_LAYOUT"


def test_word_uses_same_frozen_title_chart_width_and_alignment(db_session, project, context_factory, report_analysis, monkeypatch):
    base = generated(db_session, project, context_factory, report_analysis)
    chart = next(block for block in base.content["blocks"] if block["type"] == "chart")
    current = revise(db_session, project, context_factory, base, [
        {"op": "set_title", "text": "指定版本标题"},
        {"op": "set_chart_layout", "block_id": chart["id"], "layout": {"width_percent": 50, "alignment": "right"}},
    ])

    def render(option, path):
        import base64
        path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="))

    monkeypatch.setattr(report_service, "render_chart", render)
    document = Document(BytesIO(report_service.build_docx(current.content, current.number)))
    assert document.paragraphs[0].text == "指定版本标题"
    assert not document.styles["Title"]._element.xpath("./w:pPr/w:pBdr")
    assert round(document.inline_shapes[0].width.inches, 2) == 3.25
    assert round(document.inline_shapes[1].width.inches, 2) == 6.5
    image_paragraph = next(paragraph for paragraph in document.paragraphs if paragraph._p.xpath(".//w:drawing"))
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    assert image_paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_full_semantic_report_contains_real_classification_matrix_and_readable_sources(db_session, project, context_factory, report_analysis):
    application_id = report_analysis.patent_ids[0]
    evidence = db_session.scalar(select(Evidence).where(Evidence.analysis_id == report_analysis.id))
    evidence.content = {**evidence.content, "application_number": "PCT/TEST/1", "semantic_evidence": [
        {"theme_id": "method-visual", "quote": "uses a visual sensor for landing"}]}
    direction = {"id": "direction-visual", "title": "视觉降落", "priority": "P1", "description": "待核查候选方向",
                 "method_theme_id": "method-visual", "scenario_theme_id": "scenario-landing", "support_count": 1,
                 "supporting_application_ids": [application_id], "conditions": ["逐条核对原文"], "risks": [],
                 "phases": [{"id": "phase-review", "title": "原文核查", "actions": [{"id": "action-verify", "title": "检查权利要求"}]}]}
    report_analysis.views = {
        "technology": {"status": "ready", "data": {"classification": "semantic", "coverage": {
            "processed_count": 1, "classified_count": 1, "review_count": 0}, "taxonomy": {
                "elements": [{"id": "element-visual", "label": "视觉传感器", "count": 1}],
                "methods": [{"id": "method-visual", "label": "视觉识别", "count": 1}],
                "scenarios": [{"id": "scenario-landing", "label": "降落", "count": 1}]}}},
        "landscape": {"status": "ready", "data": {"subject_matrix": [{"applicant": "Applicant A", "theme_id": "method-visual", "count": 1}]}},
        "opportunities": {"status": "ready", "data": {"items": [direction]}},
        "roadmap": {"status": "ready", "data": {"items": [direction]}},
        "geography": {"status": "ready", "data": {"country": [{"country": "CN", "count": 1}]}},
    }
    db_session.flush()
    version = generated(db_session, project, context_factory, report_analysis)
    blocks = version.content["blocks"]
    assert any(block.get("text") == "全量语义分类与技术结构" for block in blocks)
    assert any(block.get("rows") == [["Applicant A", "视觉识别", "1"]] for block in blocks)
    assert any(block.get("section") == "roadmap" for block in blocks)
    source = next(block for block in blocks if block.get("evidence_id") == evidence.id)
    assert "PCT/TEST/1" in source["text"] and "uses a visual sensor for landing" in source["text"]
    assert "第 2 行" in source["text"] and source["application_id"] == application_id
    assert any(block.get("text") == "申请人来源国家" for block in blocks)
    assert not any("当前固定分析未形成语义分类" in block.get("text", "") for block in blocks)


def test_committed_generation_gap_reuses_ancestor_artifact_without_repeating_snapshot(db_session, project, context_factory, report_analysis, monkeypatch):
    original = context_factory(project, "generate_report", {"analysis_id": report_analysis.id})
    saved = report_service.generate_report(original)
    # Simulate a process failure after artifact commit but before run completion.
    db_session.get(ResearchRun, original.run_id).status = "failed"
    retry = context_factory(project, "generate_report", {"analysis_id": report_analysis.id})
    db_session.get(ResearchRun, retry.run_id).retry_of_id = original.run_id
    db_session.flush()
    monkeypatch.setattr(report_service, "_load_snapshot", lambda ctx: pytest.fail("Committed generation must not reload or regenerate"))
    assert report_service.generate_report(retry)["report_version_id"] == saved["report_version_id"]
    db_session.expire_all()
    assert db_session.get(Report, saved["report_id"]).version_counter == 1


def test_committed_revision_gap_ignores_stale_cas_and_preserves_current_branch(db_session, project, context_factory, report_analysis):
    base = generated(db_session, project, context_factory, report_analysis)
    payload = {"report_id": base.report_id, "base_version_id": base.id, "expected_current_version_id": base.id,
               "operations": [{"op": "set_title", "text": "Already saved revision"}]}
    original = context_factory(project, "revise_report", payload)
    saved = report_service.revise_report(original)
    db_session.expire_all()
    report = db_session.get(Report, base.report_id)
    report.current_version_id = base.id
    db_session.flush()
    branch = revise(db_session, project, context_factory, base, [{"op": "set_title", "text": "Current branch"}])
    intermediate = context_factory(project, "revise_report", payload)
    db_session.get(ResearchRun, intermediate.run_id).retry_of_id = original.run_id
    retry = context_factory(project, "revise_report", payload)
    db_session.get(ResearchRun, retry.run_id).retry_of_id = intermediate.run_id
    db_session.flush()
    recovered = report_service.revise_report(retry)
    assert recovered["report_version_id"] == saved["report_version_id"]
    assert recovered["current_version_id"] == branch.id
    db_session.expire_all()
    assert db_session.get(Report, base.report_id).version_counter == 3


def test_committed_export_gap_reuses_parent_file_without_rendering(db_session, project, context_factory, report_analysis, monkeypatch):
    base = generated(db_session, project, context_factory, report_analysis)
    payload = {"report_id": base.report_id, "report_version_id": base.id, "format": "docx"}
    original = context_factory(project, "export", payload)

    def build(*args, **kwargs):
        target = BytesIO()
        Document().save(target)
        return target.getvalue()

    monkeypatch.setattr(report_service, "build_docx", build)
    saved = report_service.export_report(original)
    retry = context_factory(project, "export", payload)
    db_session.get(ResearchRun, retry.run_id).retry_of_id = original.run_id
    db_session.flush()
    monkeypatch.setattr(report_service, "build_docx", lambda *a, **k: pytest.fail("Committed export must not render again"))
    assert report_service.export_report(retry) == saved
