from io import BytesIO

from docx import Document
from sqlalchemy import select

from app import jobs
from app.models import Chart, Message
from test_api import api_dataset, client, confirm, run


def document_bytes():
    document = Document()
    document.add_paragraph("研究需求：比较无人机自主降落的方法。")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "仅用于需求理解，不计为专利。"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_material_upload_attachment_only_and_project_isolation(client, api_dataset, monkeypatch):
    project_id = api_dataset["project"]["id"]
    path = f"/api/projects/{project_id}/materials"
    uploaded = client.post(path, files={"file": ("requirement.docx", document_bytes())})
    assert uploaded.status_code == 202, uploaded.text
    material_id, task_id = uploaded.json()["material_id"], uploaded.json()["run_id"]
    draft = {"initial_message": {"attachments": [material_id]}}
    assert client.post(f"/api/projects/{project_id}/conversations", json=draft).status_code == 409
    jobs.execute_run(task_id)
    material = client.get(f"/api/materials/{material_id}").json()
    assert material["status"] == "ready" and len(material["blocks"]) == 2
    assert material["file_name"] == "requirement.docx"
    assert "page" not in material["blocks"][0]["location"]
    assert client.get(path).json()["items"][0]["id"] == material_id
    monkeypatch.setattr(jobs, "_chat", lambda ctx: {"attachment_test": True})
    sent = client.post(f"/api/projects/{project_id}/conversations", json=draft)
    assert sent.status_code == 202, sent.text
    jobs.execute_run(sent.json()["run_id"])
    rows = client.get(f"/api/conversations/{sent.json()['conversation_id']}/messages").json()["items"]
    assert rows[0]["attachments"] == [material_id]
    assert rows[0]["content"] == ""
    assert rows[1]["cards"][0]["attachment_test"]
    other = client.post("/api/projects", json={"name": "其他项目"}).json()
    assert client.post(f"/api/projects/{other['id']}/conversations", json=draft).status_code == 404
    assert client.post(f"/api/projects/{project_id}/conversations", json={"initial_message": {"content": "   "}}).status_code == 422


def test_chart_history_routes_and_historical_drilldown(client, api_dataset, db_session):
    project = api_dataset["project"]
    confirm(client, project)
    analyzed = run(client, project["id"], "analyze")
    chart = db_session.scalar(select(Chart).where(Chart.analysis_id == analyzed["result"]["analysis_id"], Chart.kind == "yearly"))
    original = client.get(f"/api/charts/{chart.id}").json()
    revision = client.post(f"/api/charts/{chart.id}/revisions", json={
        "base_version_id": original["version_id"], "expected_current_version_id": original["version_id"],
        "mode": "data", "filters": {"start": 2024}})
    assert revision.status_code == 202, revision.text
    jobs.execute_run(revision.json()["run_id"])
    state = client.get(f"/api/runs/{revision.json()['run_id']}").json()
    assert state["status"] == "succeeded", state
    versions = client.get(f"/api/charts/{chart.id}/versions").json()
    assert len(versions["items"]) == 2
    assert versions["current_version_id"] != original["version_id"]
    old = client.get(f"/api/charts/{chart.id}", params={"version_id": original["version_id"]}).json()
    assert old["data"] == original["data"]
    found = client.get(f"/api/projects/{project['id']}/patents", params={"chart_id": chart.id,
        "chart_version_id": original["version_id"], "dimension_value": "2020"})
    assert found.status_code == 200, found.text
    assert found.json()["total"] == 1
    adopted = client.put(f"/api/charts/{chart.id}/current-version", json={"version_id": original["version_id"],
        "expected_current_version_id": versions["current_version_id"]})
    assert adopted.status_code == 200, adopted.text
    assert client.get(f"/api/charts/{chart.id}").json()["version_id"] == original["version_id"]


def test_operations_and_legacy_report_revision_are_mutually_exclusive(client, api_dataset):
    body = {"base_version_id": api_dataset["project"]["id"], "expected_current_version_id": api_dataset["project"]["id"],
            "block_id": "block", "replacement_text": "text", "operations": [{"op": "set_title", "text": "标题"}]}
    response = client.post(f"/api/reports/{api_dataset['project']['id']}/revisions", json=body)
    assert response.status_code == 422


def test_material_same_name_preserves_distinct_file_versions(client, api_dataset):
    path = f"/api/projects/{api_dataset['project']['id']}/materials"
    first = client.post(path, files={"file": ("brief.docx", document_bytes())}).json()
    document = Document()
    document.add_paragraph("更新后的研究目标")
    content = BytesIO()
    document.save(content)
    second = client.post(path, files={"file": ("brief.docx", content.getvalue())}).json()
    first_value = client.get(f"/api/materials/{first['material_id']}").json()
    second_value = client.get(f"/api/materials/{second['material_id']}").json()
    assert first_value["version"] == 1 and second_value["version"] == 2
    assert first_value["file_id"] != second_value["file_id"]
    assert first_value["sha256"] != second_value["sha256"]
